"""Excel export — consumes the compute core's resolved values (ENGINE_CONTRACT §11).

This module imports openpyxl; the pure core (engine/compute.py) never does. The
workbook is *formula-driven*: inputs are blue hard-coded cells, computed cells are
live formulas over them, totals are =SUM(...) over their components, and every
divided cell is wrapped =IF(denom=0,"",num/denom) so a LibreOffice recalc never
yields #DIV/0!. Overridden components become hard-coded orange cells carrying the
exact audit comment; the enclosing total stays a formula so the sheet foots.

The three-way golden gate (test_golden_gate.py) recalcs this workbook in
LibreOffice and asserts every cell equals the engine's resolved value within
$0.50 / 0.1pp with zero formula errors.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from engine import dag
from engine.compute import EXPENSE_ORDER
from engine.money import D, round_dollars

BLUE = Font(name="Arial", color="0000FF", size=10)
BLACK = Font(name="Arial", color="000000", size=10)
BOLD = Font(name="Arial", bold=True, size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
OVERRIDE_FILL = PatternFill("solid", fgColor="F4C9A3")  # orange (overridden)
TOT_FILL = PatternFill("solid", fgColor="EAF1F4")
USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%'
NUM = '#,##0'

_OVERRIDE_COMMENT = (
    "Manual override — calc was ={formula} → ${calc:,.0f}; "
    "set to ${value:,.0f} on {date}")


def _f(value):
    """Decimal -> float for openpyxl numeric cells (display rounding done by Excel)."""
    return float(D(str(value)))


def build_xlsx(result, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Projection"
    monthly_ws = wb.create_sheet("Monthly Cash Flow")

    inputs = result.inputs
    a = result.assumptions
    overrides = result.overrides_applied
    pinned = {}
    for key in overrides:
        node, year = dag.parse_key(key)
        pinned[(node, year)] = overrides[key]

    # ---------- INPUTS block (blue) ----------------------------------------
    ws["A1"] = f"{inputs.org_name} — 3-Year Thrift Store Projection"
    ws["A1"].font = Font(name="Arial", bold=True, size=14)

    cell_ref = {}
    row = 3
    ws[f"A{row}"] = "KEY INPUTS"
    ws[f"A{row}"].font = BOLD
    row += 1
    input_specs = [
        ("households", _f(result.trade_area.households), NUM),
        ("donations_per_hh", a.donations_per_hh, NUM),
        ("share_y1", _f(inputs.share_y1), PCT),
        ("share_y2", _f(inputs.share_y2), PCT),
        ("share_y3", _f(inputs.share_y3), PCT),
        ("donation_value", _f(inputs.donation_value), USD),
        ("basket", _f(a.basket), USD),
        ("salvage_pct", _f(a.salvage_pct), PCT),
        ("card_pct", _f(a.card_pct), PCT),
        ("labor_pct", _f(a.labor_pct), PCT),
        ("payroll_tax_pct", _f(a.payroll_tax_pct), PCT),
        ("marketing_pct", _f(a.marketing_pct), PCT),
        ("supplies_pct", _f(a.supplies_pct), PCT),
        ("occupancy_y1", _f(result.occupancy[1]), USD),
        ("escalator", _f(a.escalator), PCT),
    ]
    for name, value, fmt in input_specs:
        ws[f"A{row}"] = name
        ws[f"A{row}"].font = BLACK
        c = ws[f"B{row}"]
        c.value = value
        c.font = BLUE
        c.fill = INPUT_FILL
        c.number_format = fmt
        cell_ref[name] = f"$B${row}"
        row += 1
    # fixed annual lines as blue inputs
    for label, value in a.fixed_annual.items():
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = BLACK
        c = ws[f"B{row}"]
        c.value = float(value)
        c.font = BLUE
        c.fill = INPUT_FILL
        c.number_format = USD
        cell_ref[label] = f"$B${row}"
        row += 1

    # ---------- 3-YEAR PROJECTION ------------------------------------------
    row += 1
    ws[f"A{row}"] = "3-YEAR PROJECTION"
    ws[f"A{row}"].font = BOLD
    row += 1
    cols = ["B", "C", "D"]
    for i, y in enumerate(("Year 1", "Year 2", "Year 3")):
        ws.cell(row=row, column=2 + i, value=y).font = BOLD
    row += 1

    rowmap = {}

    def line(label, formulas, fmt=USD, bold=False, fill=None):
        nonlocal row
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = BOLD if bold else BLACK
        for i, val in enumerate(formulas):
            c = ws.cell(row=row, column=2 + i, value=val)
            c.font = BOLD if bold else BLACK
            c.number_format = fmt
            c.alignment = Alignment(horizontal="right")
            if fill:
                c.fill = fill
        rowmap[label] = row
        row += 1

    sh = cell_ref
    # share row — three independent input cells (NO average).
    line("Donor market share",
         [f"={sh['share_y1']}", f"={sh['share_y2']}", f"={sh['share_y3']}"], PCT)
    sr = rowmap["Donor market share"]
    line("Donating households",
         [f"={c}{sr}*{sh['households']}" for c in cols], NUM)
    dh = rowmap["Donating households"]
    line("Annual donations",
         [f"={c}{dh}*{sh['donations_per_hh']}" for c in cols], NUM)
    dn = rowmap["Annual donations"]
    line("Retail Sales", [f"={c}{dn}*{sh['donation_value']}" for c in cols], USD)
    rt = rowmap["Retail Sales"]
    line("Salvage Sales", [f"={c}{rt}*{sh['salvage_pct']}" for c in cols], USD)
    sv = rowmap["Salvage Sales"]
    line("Total Revenue", [f"={c}{rt}+{c}{sv}" for c in cols], USD, bold=True)
    tr = rowmap["Total Revenue"]
    # customers = retail / basket  (guarded)
    line("Customers",
         [f'=IF({sh["basket"]}=0,"",{c}{rt}/{sh["basket"]})' for c in cols], NUM)

    # ---- expense lines (live formulas, or orange override) ----
    exp_rows = []
    for label in EXPENSE_ORDER:
        if label not in result.annual[1].expenses:
            continue
        formulas = []
        for i, c in enumerate(cols):
            year = i + 1
            calc_formula = _expense_formula(label, c, tr, sh, rowmap)
            if (label, year) in pinned:
                formulas.append(None)  # placeholder; set as hard value below
            else:
                formulas.append(f"={calc_formula}")
        line(label, [f if f is not None else 0 for f in formulas], USD)
        r = rowmap[label]
        exp_rows.append(r)
        # apply overrides: hard-coded orange + exact comment
        for i, c in enumerate(cols):
            year = i + 1
            if (label, year) in pinned:
                pin = pinned[(label, year)]
                # recompute the *pre-override* calc value for the comment
                pre = _precalc_expense(label, year, result)
                cell = ws.cell(row=r, column=2 + i)
                cell.value = _f(pin["value"])
                cell.fill = OVERRIDE_FILL
                cell.number_format = USD
                cell.comment = Comment(_OVERRIDE_COMMENT.format(
                    formula=_expense_formula(label, c, tr, sh, rowmap).lstrip("="),
                    calc=float(pre), value=float(D(pin["value"])), date=pin["ts"]), "engine")

    # totals = SUM over component cells (foots even with overrides)
    first, last = min(exp_rows), max(exp_rows)
    line("Total Expenses", [f"=SUM({c}{first}:{c}{last})" for c in cols], USD, bold=True)
    te = rowmap["Total Expenses"]
    line("Net profit", [f"={c}{tr}-{c}{te}" for c in cols], USD, bold=True, fill=TOT_FILL)
    npr = rowmap["Net profit"]
    # net margin = net/totrev (guarded)
    line("Net margin",
         [f'=IF({c}{tr}=0,"",{c}{npr}/{c}{tr})' for c in cols], PCT)

    # ---------- STARTUP + payback (guarded) --------------------------------
    row += 1
    ws[f"A{row}"] = "ONE-TIME STARTUP COSTS"
    ws[f"A{row}"].font = BOLD
    row += 1
    su_first = row
    for label, value in result.startup.items():
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = BLACK
        c = ws[f"B{row}"]
        if label == "Lease Deposit + First Month Rent":
            c.value = f"={sh['occupancy_y1']}/12*3"
        else:
            c.value = float(value)
            c.fill = INPUT_FILL
            c.font = BLUE
        c.number_format = USD
        row += 1
    su_last = row - 1
    ws[f"A{row}"] = "Total Startup Costs"
    ws[f"A{row}"].font = BOLD
    stc = ws[f"B{row}"]
    stc.value = f"=SUM(B{su_first}:B{su_last})"
    stc.number_format = USD
    stc.fill = TOT_FILL
    startup_total_row = row
    row += 1
    ws[f"A{row}"] = "Simple payback (years)"
    ws[f"A{row}"].font = BLACK
    pc = ws[f"B{row}"]
    pc.value = f'=IF(B{npr}<=0,"",B{startup_total_row}/B{npr})'
    pc.number_format = '0.0'
    row += 1

    # ---------- MILESTONES (grand opening / break-even / payback) -----------
    # The three timeline milestones (§8 / #117), rendered as the SAME server display
    # strings the view-model emits ("Month 6" / "N/A"). These are engine-resolved
    # values, not divisions — no formula error surface to guard.
    _build_milestones_block(ws, result, row)

    # ---------- MONTHLY tab -------------------------------------------------
    _build_monthly_tab(monthly_ws, result)

    ws.column_dimensions["A"].width = 34
    wb.save(path)


def _expense_formula(label, col, tr, sh, rowmap):
    """The live Excel formula for an expense component (sans leading '=')."""
    if label == "Credit card fees":
        return f"{col}{tr}*{sh['card_pct']}"
    if label == "Marketing & outreach":
        return f"{col}{tr}*{sh['marketing_pct']}"
    if label == "Payroll":
        return f"{col}{tr}*{sh['labor_pct']}"
    if label == "Payroll taxes":
        pr = rowmap.get("Payroll")
        return f"{col}{pr}*{sh['payroll_tax_pct']}"
    if label == "Supplies & processing":
        return f"{col}{tr}*{sh['supplies_pct']}"
    if label == "Rent":
        # occupancy_y1 * (1+escalator)^(year-1); year derived from column.
        yexp = {"B": 0, "C": 1, "D": 2}[col]
        return f"{sh['occupancy_y1']}*(1+{sh['escalator']})^{yexp}"
    # fixed annual line
    return f"{sh[label]}"


def _precalc_expense(label, year, result):
    """The pre-override calculated value (for the override comment string)."""
    # Recompute the component from the resolved (non-overridden) parents.
    from engine.compute import compute_annual
    years, _su, _st, _occ = compute_annual(
        result.inputs, result.assumptions, None, result.trade_area, {})
    return round_dollars(years[year].expenses[label])


def _build_milestones_block(ws, result, row):
    """Grand opening / break-even / payback as the SAME server display strings the
    view-model emits ("Month 6" / "N/A"), so the workbook and the app never disagree.

    Reuses ``engine.viewmodel._milestone`` (the single formatting authority) rather
    than re-deriving the strings here. These are engine-resolved values, not Excel
    divisions, so there is nothing to guard against #DIV/0!.
    """
    from engine.viewmodel import _milestone

    dec = result.decision
    g = dec.grand_opening_month
    entries = [
        ("Grand opening", _milestone(g, result.flags, "grand_opening_month")),
        ("Break-even", _milestone(dec.break_even_month, result.flags, "break_even_month")),
        ("Payback", _milestone(dec.payback_month, result.flags, "payback_month",
                               opening_month=g)),
    ]
    ws[f"A{row}"] = "MILESTONES"
    ws[f"A{row}"].font = BOLD
    row += 1
    for label, milestone in entries:
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = BLACK
        ws[f"B{row}"] = milestone["display"]
        row += 1


def _build_monthly_tab(ws, result):
    """Monthly cash flow over the uniform project-month timeline 1..N (Spec §6).

    Every project month gets a P&L row. Pre-opening months (1..G-1) carry no
    operating revenue/expense — their scheduled one-time costs land on their own
    month #, and Net = revenue - opex - scheduled costs is a live formula, so a
    pre-opening month reads as a real monthly loss (negative net). Operating
    months (G..N) carry the ramp plus any scheduled costs on the same month. Net
    and Cumulative cash are formulas; no division, so nothing to guard here.
    """
    g = result.inputs.opening_month_number
    ws["A1"] = "Monthly Cash Flow (project months 1..N)"
    ws["A1"].font = BOLD
    # Columns: A Month | B Revenue | C Operating expenses | D Scheduled costs
    #          | E Net | F Cumulative cash
    headers = ["Month", "Revenue", "Operating expenses", "Scheduled costs",
               "Net", "Cumulative cash"]
    for i, hdr in enumerate(headers):
        c = ws.cell(row=2, column=1 + i, value=hdr)
        c.font = BOLD
    r = 3
    first_data = r
    for m in result.monthly:
        label = f"Month {m.index}"
        if m.index == g:
            label = f"Month {m.index} — Grand opening"
        lc = ws.cell(row=r, column=1, value=label)
        if m.index == g:
            lc.font = BOLD
        ws.cell(row=r, column=2, value=_f(m.revenue)).number_format = USD
        ws.cell(row=r, column=3, value=_f(m.expenses)).number_format = USD
        ws.cell(row=r, column=4, value=_f(m.startup_outflow)).number_format = USD
        # Net = revenue - operating expenses - scheduled costs (live formula).
        ws.cell(row=r, column=5, value=f"=B{r}-C{r}-D{r}").number_format = USD
        # cumulative cash = running sum of Net (Net already nets out scheduled costs)
        if r == first_data:
            ws.cell(row=r, column=6, value=f"=E{r}").number_format = USD
        else:
            ws.cell(row=r, column=6, value=f"=F{r-1}+E{r}").number_format = USD
        r += 1
    ws.column_dimensions["A"].width = 24
