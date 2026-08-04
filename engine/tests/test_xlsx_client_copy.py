"""RED first: the client-copy ``.xlsx`` variant (``?variant=client``, Model v3
Phase 3, Task 3.1).

Reuses the SAME engine core (``compute()``) as the internal ``.xlsx`` — this test
file only exercises the NEW ``build_client_copy_xlsx`` renderer, never the
existing ``build_xlsx`` (that renderer's own tests, ``engine/tests/test_xlsx.py``,
must stay green/untouched — proven by the broader sweep in the phase's
verification proof set, not re-asserted here).

Four things a board-facing client copy must never do:
  (a) carry any KPI / sensitivity / guardrails / decision-metrics block — just
      the Projection sheet's revenue/expense rows, the Startup Costs sheet, and
      Planning Notes;
  (b) show a zero-dollar startup line as a row (dash/zero rows read as noise);
  (c) omit the exact, non-paraphrased planning-estimates disclaimer; and
  (d) leak internal vocabulary — a "§" section citation or the word "override"
      (case-insensitive) — anywhere in any cell value or comment, even when a
      pinned/overridden cell is present in the fixture.
"""
import openpyxl

from engine.compute import Assumptions, DemandDrivers, Inputs, TradeArea, compute
from engine.model import CLIENT_COPY_DISCLAIMER
from engine.xlsx import build_client_copy_xlsx

_OVERRIDE = {
    "Payroll#1": {
        "value": "110000",
        "note": "2 FTE managers + PT staff, benefits (justification)",
        "ts": "2026-08-04T00:00:00Z",
    }
}


def _cost_schedule():
    return (
        {"category": "Equipment & Fixtures", "month": 1, "amount": "45000",
         "treatment": "included"},
        # A zero-dollar included line — must never render as a row (b).
        {"category": "Signage Permit Fee", "month": 1, "amount": "0",
         "treatment": "included"},
        {"category": "Thrift Store Consultant", "month": 1, "amount": "100000",
         "treatment": "separate",
         "note": "Handled under a separate consulting agreement."},
        {"category": "Contingency Reserve", "month": 1, "amount": "15000",
         "treatment": "contingent"},
        # A zero-dollar contingent line — must also never render as a row (b).
        {"category": "Unexpected Permit Delay", "month": 1, "amount": "0",
         "treatment": "contingent"},
    )


def _inputs():
    return Inputs(
        org_name="House of Hope", city="Longview, Texas", opening_status="pre_launch",
        share_y1="0.06", share_y2="0.08", share_y3="0.10", share_y4="0.11", share_y5="0.12",
        donation_value="74", building_sqft="10000", salesfloor_pct="0.70",
        occupancy_annual="120000", lease_type="nnn", source_of_funds="300000",
        cost_schedule=_cost_schedule(),
    )


def _build(tmp_path, overrides=None):
    res = compute(_inputs(), Assumptions(), DemandDrivers(),
                  TradeArea(basis="city", households="31798"),
                  overrides=overrides or {})
    path = tmp_path / "client-copy.xlsx"
    build_client_copy_xlsx(res, str(path))
    return res, path


def _all_cell_texts(wb):
    """Every non-empty cell value + every comment's text, across every sheet."""
    texts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    texts.append(str(cell.value))
                if cell.comment is not None:
                    texts.append(str(cell.comment.text))
    return texts


# --------------------------------------------------------------------- (a)
def test_no_kpi_sensitivity_or_guardrails_block(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)

    banned_sheet_words = ("kpi", "sensitivity", "guardrail", "milestone")
    for name in wb.sheetnames:
        lowered = name.lower()
        assert not any(w in lowered for w in banned_sheet_words), name

    blob = " ".join(_all_cell_texts(wb)).lower()
    for banned in ("kpi", "sensitivity", "guardrail", "milestone",
                   "break-even", "payback", "decision"):
        assert banned not in blob, banned


def test_only_three_sheets_projection_startup_notes(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    assert set(wb.sheetnames) == {"Projection", "Startup Costs", "Planning Notes"}


# --------------------------------------------------------------------- (b)
def test_zero_value_startup_line_omitted(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Startup Costs"]
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Signage Permit Fee" not in texts
    assert "Unexpected Permit Delay" not in texts
    # sanity: the non-zero rows in the SAME categories DID render, proving this is
    # a value-filter, not an accidental category-name filter.
    assert "Equipment & Fixtures" in texts
    assert "Contingency Reserve" in texts


def test_separate_and_contingent_blocks_are_plain_language_labeled(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Startup Costs"]
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Included in capital required" in texts
    assert "Funded separately" in texts
    assert "Contingent" in texts
    # No internal jargon anywhere in these labels.
    for t in texts:
        assert "treatment" not in t.lower()


# --------------------------------------------------------------------- (c)
def test_planning_notes_sheet_has_exact_disclaimer(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    assert "Planning Notes" in wb.sheetnames
    ws = wb["Planning Notes"]
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert CLIENT_COPY_DISCLAIMER in texts


def test_planning_notes_covers_the_major_assumption_groups(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    blob = " ".join(str(c.value) for row in wb["Planning Notes"].iter_rows()
                     for c in row if c.value is not None)
    assert "market share" in blob.lower()
    assert "payroll" in blob.lower()
    assert "rent" in blob.lower()
    assert "card" in blob.lower()


# --------------------------------------------------------------------- (d)
def test_no_section_citation_or_override_word_anywhere(tmp_path):
    _res, path = _build(tmp_path, overrides=_OVERRIDE)
    wb = openpyxl.load_workbook(path)
    texts = _all_cell_texts(wb)
    blob = " ".join(texts)
    assert "§" not in blob
    assert "override" not in blob.lower()
    # sanity: the pinned cell really landed with an adjustment comment, so the
    # assertion above is meaningfully exercised rather than trivially true.
    assert any("Manually adjusted" in t for t in texts)
