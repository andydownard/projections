"""Pin the `.xlsx` ANNUAL formula cascade (ENGINE_CONTRACT §11.8).

The annual "Projection" sheet is built as a *live* spreadsheet: every blue input
cell drives the driver chain, the expense lines, the totals, Net income and Net
margin as Excel FORMULAS, so a board member who opens the workbook and edits an
assumption sees the whole P&L recompute. Two test layers pin that contract:

  1. AUDIT (always runs, no soffice): every DERIVED annual cell on the Projection
     sheet is an Excel formula (the openpyxl value is a ``str`` starting with
     "="), never a hard-coded number. This is the structural regression guard:
     a future change can never silently bake a derived annual value into a literal
     and quietly break the cascade.

  2. CASCADE (LibreOffice recalc, gated on ``soffice`` EXACTLY like
     test_golden_gate.py: runs in CI where soffice is pinned, skips locally):
     edit ONE blue input cell, recalc the workbook, and assert the downstream
     cells moved to the new hand-computed values. This proves "edit one cell ->
     it cascades" end to end through LibreOffice, not just in openpyxl strings.

OUT OF SCOPE (deliberate, separate decision): the Monthly tab per-month
revenue/opex cells and the break-even / payback MILESTONES are server VALUE
snapshots, not live formulas; they are not asserted here.
"""
import pathlib
import shutil
import sys
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from engine.compute import (
    Assumptions,
    DemandDrivers,
    EXPENSE_ORDER,
    Inputs,
    TradeArea,
    cash_margin,
    cash_net,
    compute,
)
from engine.xlsx import build_xlsx

# The LibreOffice recalc harness (soffice + a Basic recalc macro) lives in the
# reference skill scripts, same path bootstrap as test_golden_gate.py.
_SCRIPTS = (
    pathlib.Path(__file__).resolve().parents[2]
    / "reference/thrift-store-projections-skill/scripts"
)
sys.path.insert(0, str(_SCRIPTS))

# IDENTICAL gate to test_golden_gate.py: the cascade test recalcs in LibreOffice,
# so it runs where soffice is installed (CI pins it) and skips otherwise. The AUDIT
# test below carries NO such marker: it is a pure openpyxl read and must run
# everywhere, including locally without soffice.
_needs_soffice = pytest.mark.skipif(
    shutil.which("soffice") is None, reason="LibreOffice (soffice) not installed"
)

_USD_TOL = 0.5      # PRD §18 dollar/count tolerance
_PCT_TOL = 0.001    # PRD §18 percent (0.1pp) tolerance

# Canonical House of Hope constants (engine/tests/conftest.py / ENGINE_CONTRACT §10),
# reused as the hand-computation basis for the cascade expectations.
_H = Decimal("31798")
_DONATIONS_PER_HH = Decimal("3")
_DONATION_VALUE = Decimal("74")
_SALVAGE_PCT = Decimal("0.008")
_LABOR_PCT = Decimal("0.60")
_BASE_SHARE_Y1 = Decimal("0.06")


def _result(share_y1="0.06", basket=None):
    """The canonical 60%-labor House of Hope result, optionally with one input edited."""
    assumptions = (
        Assumptions(labor_pct="0.60")
        if basket is None
        else Assumptions(labor_pct="0.60", basket=basket)
    )
    inputs = Inputs(
        org_name="House of Hope", city="Longview, Texas", opening_status="pre_launch",
        share_y1=share_y1, share_y2="0.08", share_y3="0.10", share_y4="0.11", share_y5="0.12",
        donation_value="74", building_sqft="10000", salesfloor_pct="0.70",
        occupancy_annual="120000", lease_type="nnn", source_of_funds="300000",
    )
    return compute(inputs, assumptions, DemandDrivers(),
                   TradeArea(basis="city", households="31798"))


def _label_row(ws):
    """Map a column-A string label -> its row index on the Projection sheet."""
    return {
        ws.cell(row=r, column=1).value: r
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=1).value, str)
    }


def _is_formula(v):
    return isinstance(v, str) and v.startswith("=")


# --------------------------------------------------------------------------- #
# 1. AUDIT: every derived annual cell is a live formula (runs everywhere)
# --------------------------------------------------------------------------- #
def test_every_derived_annual_cell_is_a_formula(tmp_path):
    """REGRESSION GUARD (no soffice needed): no derived annual cell may be a literal.

    Build the canonical workbook (NO overrides) and assert every derived cell on
    the Projection sheet (the full driver chain, every expense line, the totals,
    Net income across all five years, Net margin, and Total Startup Costs) is a
    live Excel formula, not a hard-coded number. If any derived annual cell ever
    becomes a literal the live cascade is broken and a board member's edit would
    silently fail to propagate; this test fails closed before that can ship.
    """
    res = _result()
    path = tmp_path / "audit.xlsx"
    build_xlsx(res, str(path))
    ws = load_workbook(path)["Projection"]
    rows = _label_row(ws)

    # The driver chain + revenue + customers + every expense line + totals + the
    # cash Net income (all 5 years) + Net margin. Year 1..5 == columns B..F (2..6).
    five_col_derived = [
        "Donor market share",
        "Donating households",
        "Annual donations",
        "Retail Sales",
        "Salvage Sales",
        "Total Revenue",
        "Customers",
        *[label for label in EXPENSE_ORDER if label in res.annual[1].expenses],
        "Total Expenses",
        "Net income",
        "Net margin",
    ]
    for label in five_col_derived:
        assert label in rows, f"Projection sheet is missing the {label!r} row"
        r = rows[label]
        for col in range(2, 7):  # B..F = Year 1..5
            v = ws.cell(row=r, column=col).value
            assert _is_formula(v), (
                f"{label!r} (year column {col}) is a hard-coded literal ({v!r}); "
                f"every derived annual cell MUST be a live Excel formula so editing "
                f"a blue input cascades through it."
            )

    # Total Startup Costs is a single =SUM(...) over the startup component cells (col B).
    assert "Total Startup Costs" in rows, "missing Total Startup Costs row"
    stc = ws.cell(row=rows["Total Startup Costs"], column=2).value
    assert _is_formula(stc) and stc.startswith("=SUM("), (
        f"Total Startup Costs must be a =SUM(...) formula, got {stc!r}"
    )

    # Belt-and-braces: the headline totals are genuinely formulas referencing OTHER
    # cells (a literal that merely starts with '=' could never reference a cell).
    assert "+" in ws.cell(row=rows["Total Revenue"], column=2).value      # retail + salvage
    assert ws.cell(row=rows["Total Expenses"], column=2).value.startswith("=SUM(")
    assert "-" in ws.cell(row=rows["Net income"], column=2).value         # rev - exp - startup


# --------------------------------------------------------------------------- #
# 2. CASCADE: edit one blue input, recalc in LibreOffice, downstream moves
# --------------------------------------------------------------------------- #
def _recalc(path):
    """Recalc ``path`` in LibreOffice and assert it computed with zero formula errors."""
    import recalc as recalc_harness

    report = recalc_harness.recalc(str(path), timeout=90)
    assert "error" not in report, report
    assert report["status"] == "success", report
    assert report["total_errors"] == 0, report.get("error_summary")


def _year1(path, label):
    """The recalc'd Year-1 (column B) numeric value of a Projection row, post-recalc."""
    wb = load_workbook(str(path), data_only=True)
    ws = wb["Projection"]
    val = ws.cell(row=_label_row(ws)[label], column=2).value
    wb.close()
    return val


def _edit_input(path, input_label, new_value):
    """Edit the blue input cell (column B of the INPUTS row whose column-A == label).

    Loads WITHOUT ``data_only`` so every downstream formula string is preserved; only
    the one input literal changes, so the next recalc re-evaluates the whole cascade.
    """
    wb = load_workbook(str(path))
    ws = wb["Projection"]
    ws.cell(row=_label_row(ws)[input_label], column=2).value = new_value
    wb.save(str(path))


@_needs_soffice
def test_edit_share_y1_cascades_to_revenue_payroll_net_margin(tmp_path):
    """Edit the ``share_y1`` blue input -> Total Revenue, Payroll, Net income and Net
    margin all move to the new values after a LibreOffice recalc.

    ``share_y1`` sits at the head of the driver chain (share -> donating households ->
    donations -> retail -> total revenue), and Payroll/Net/margin are all revenue-driven,
    so this is the strongest possible "one edit cascades through the whole P&L" proof.
    """
    base = _result()

    # (a) UNEDITED recalc == the server's resolved values (sanity: the workbook the
    #     board first opens already foots to the engine).
    unedited = tmp_path / "share_unedited.xlsx"
    build_xlsx(base, str(unedited))
    _recalc(unedited)
    assert _year1(unedited, "Total Revenue") == pytest.approx(
        float(base.annual[1].total_rev), abs=_USD_TOL)
    assert _year1(unedited, "Payroll") == pytest.approx(
        float(base.annual[1].expenses["Payroll"]), abs=_USD_TOL)
    assert _year1(unedited, "Net income") == pytest.approx(
        float(cash_net(base, 1)), abs=_USD_TOL)
    assert _year1(unedited, "Net margin") == pytest.approx(
        float(cash_margin(base, 1)), abs=_PCT_TOL)

    # (b) EDIT share_y1 0.06 -> 0.12 and recalc the edited workbook.
    new_share = Decimal("0.12")
    edited = tmp_path / "share_edited.xlsx"
    build_xlsx(base, str(edited))
    _edit_input(edited, "share_y1", float(new_share))
    _recalc(edited)

    # Hand-computed closed forms for the two tractable cells (independent of the engine):
    #   total revenue = H x share x donations/hh x donation_value x (1 + salvage_pct)
    #   payroll       = total revenue x labor_pct
    rev_hand = float(
        _H * new_share * _DONATIONS_PER_HH * _DONATION_VALUE * (Decimal("1") + _SALVAGE_PCT))
    payroll_hand = float(Decimal(str(rev_hand)) * _LABOR_PCT)
    # Engine recompute with the SAME edit = the expectation for Net income / Net margin
    # (LibreOffice formula evaluation vs the Python engine is an independent cross-check,
    # exactly as the three-way golden gate reconciles them).
    exp = _result(share_y1="0.12")

    tr = _year1(edited, "Total Revenue")
    pr = _year1(edited, "Payroll")
    ni = _year1(edited, "Net income")
    nm = _year1(edited, "Net margin")

    # Each cell MOVED away from the unedited recalc...
    assert tr != pytest.approx(float(base.annual[1].total_rev), abs=_USD_TOL)
    assert pr != pytest.approx(float(base.annual[1].expenses["Payroll"]), abs=_USD_TOL)
    assert ni != pytest.approx(float(cash_net(base, 1)), abs=_USD_TOL)
    assert nm != pytest.approx(float(cash_margin(base, 1)), abs=_PCT_TOL)
    # ...to the new hand-computed AND engine-recomputed expectation.
    assert tr == pytest.approx(rev_hand, abs=_USD_TOL)
    assert tr == pytest.approx(float(exp.annual[1].total_rev), abs=_USD_TOL)
    assert pr == pytest.approx(payroll_hand, abs=_USD_TOL)
    assert pr == pytest.approx(float(exp.annual[1].expenses["Payroll"]), abs=_USD_TOL)
    assert ni == pytest.approx(float(cash_net(exp, 1)), abs=_USD_TOL)
    assert nm == pytest.approx(float(cash_margin(exp, 1)), abs=_PCT_TOL)


@_needs_soffice
def test_edit_basket_cascades_to_customers_only(tmp_path):
    """Edit the ``basket`` blue input -> the Customers count moves (it is the divisor),
    while Total Revenue stays put.

    In this model basket value drives ONLY Customers (= retail / basket); revenue is
    donations-driven (retail = donations x donation_value), so basket is a pure
    customers divisor. Asserting Customers moves AND Total Revenue does NOT proves the
    cascade is *precise*: an edit propagates to exactly the cells it should and to no
    others.
    """
    base = _result()

    # (a) UNEDITED recalc == server resolved (sanity).
    unedited = tmp_path / "basket_unedited.xlsx"
    build_xlsx(base, str(unedited))
    _recalc(unedited)
    base_customers = _year1(unedited, "Customers")
    base_rev = _year1(unedited, "Total Revenue")
    assert base_customers == pytest.approx(float(base.annual[1].customers), abs=_USD_TOL)
    assert base_rev == pytest.approx(float(base.annual[1].total_rev), abs=_USD_TOL)

    # (b) EDIT basket $17.08 -> $25.00 and recalc.
    new_basket = Decimal("25.00")
    edited = tmp_path / "basket_edited.xlsx"
    build_xlsx(base, str(edited))
    _edit_input(edited, "basket", float(new_basket))
    _recalc(edited)

    # Hand-computed: customers = retail / basket, retail = H x share x donations/hh x value.
    retail_hand = _H * _BASE_SHARE_Y1 * _DONATIONS_PER_HH * _DONATION_VALUE
    customers_hand = float(retail_hand / new_basket)
    exp = _result(basket="25.00")

    cust = _year1(edited, "Customers")
    rev = _year1(edited, "Total Revenue")

    # Customers MOVED to the new value (basket is the divisor)...
    assert cust != pytest.approx(base_customers, abs=_USD_TOL)
    assert cust == pytest.approx(customers_hand, abs=_USD_TOL)
    assert cust == pytest.approx(float(exp.annual[1].customers), abs=_USD_TOL)
    # ...and Total Revenue is UNCHANGED: basket cascades only where the formulas wire it.
    assert rev == pytest.approx(base_rev, abs=_USD_TOL)
    assert rev == pytest.approx(float(base.annual[1].total_rev), abs=_USD_TOL)
