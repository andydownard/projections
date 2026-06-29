"""RED first: sq-ft-scaled utilities + fixtures/build-out seed (feature #245).

The example store (House of Hope) is 10,000 sq ft. Every per-sq-ft constant is
derived from that basis (utilities $22,300/yr, Equipment & Fixtures $45,000,
Build-out $50,000) and ANCHORED so that at 10,000 sq ft it reproduces today's
literal byte-for-byte. Utilities stops being a flat fixed-$ line and becomes the
SAME existing line derived as ``UTILITIES_PER_SQFT × building_sqft``; the default
cost-schedule seed scales Equipment & Fixtures / Build-out the same way.

These are the §1/§4 anchor invariants of the plan: 10k = unchanged, other sizes
move linearly.
"""
from decimal import Decimal

from engine import model
from engine.compute import (
    Assumptions,
    DemandDrivers,
    Inputs,
    TradeArea,
    compute,
)
from engine.money import D


# --------------------------------------------------------------- T1.1 constants
def test_per_sqft_constants_anchor_to_example_literals():
    sqft = Decimal(str(model.EXAMPLE_BUILDING_SQFT))
    assert sqft == Decimal("10000")
    # Each constant × the example's 10,000 sq ft reproduces today's literal exactly.
    assert D(model.UTILITIES_PER_SQFT) * sqft == Decimal("22300")
    assert D(model.FIXTURES_PER_SQFT) * sqft == Decimal("45000")
    assert D(model.BUILDOUT_PER_SQFT) * sqft == Decimal("50000")


def test_default_rent_per_sqft_is_golden_parity_12():
    # AI-off default pins occupancy at $12/sq ft so 10k resolves to today's $120k.
    assert D(model.DEFAULT_RENT_PER_SQFT) == Decimal("12.00")
    assert D(model.DEFAULT_RENT_PER_SQFT) * Decimal("10000") == Decimal("120000")


def test_utilities_removed_from_fixed_annual():
    # Utilities is no longer a flat fixed-$ line; it scales with sq ft now.
    assert "Utilities/trash service" not in model.FIXED_ANNUAL


# --------------------------------------------------- T1.2 utilities scales w/ ft
def _inputs(building_sqft="10000"):
    return Inputs(
        org_name="House of Hope", city="Longview, Texas",
        opening_status="pre_launch",
        share_y1="0.06", share_y2="0.08", share_y3="0.10", share_y4="0.11", share_y5="0.12",
        donation_value="74", building_sqft=building_sqft, salesfloor_pct="0.70",
        occupancy_annual="120000", lease_type="nnn", source_of_funds="300000",
    )


def _trade():
    return TradeArea(basis="city", households="31798")


def test_utilities_at_10k_reproduces_todays_literal():
    res = compute(_inputs("10000"), Assumptions(), DemandDrivers(), _trade())
    for y in (1, 2, 3):
        assert D(str(res.annual[y].expenses["Utilities/trash service"])) == Decimal("22300")


def test_utilities_scales_linearly_with_sqft():
    res = compute(_inputs("8000"), Assumptions(), DemandDrivers(), _trade())
    # 2.23 × 8000 = 17,840 (the Grace-Mission-class move).
    assert D(str(res.annual[1].expenses["Utilities/trash service"])) == Decimal("17840")
    res2 = compute(_inputs("20000"), Assumptions(), DemandDrivers(), _trade())
    assert D(str(res2.annual[1].expenses["Utilities/trash service"])) == Decimal("44600")


def test_utilities_remains_a_pinnable_component_line():
    # Same label, same slot, still overridable — no new DAG node.
    res = compute(
        _inputs("10000"), Assumptions(), DemandDrivers(), _trade(),
        overrides={"Utilities/trash service#1": {"value": "5000", "note": "x", "ts": "t"}},
    )
    assert D(str(res.annual[1].expenses["Utilities/trash service"])) == Decimal("5000")
    # Untouched years still derive from sq ft.
    assert D(str(res.annual[2].expenses["Utilities/trash service"])) == Decimal("22300")


def test_utilities_assumption_default_is_the_constant():
    assert D(Assumptions().utilities_per_sqft) == D(model.UTILITIES_PER_SQFT)


# ------------------------------------------------- T1.3 fixtures/build-out seed
def test_default_cost_schedule_at_10k_is_byte_identical():
    # No-arg + explicit 10,000 both reproduce today's frozen literals exactly.
    frozen = [
        {"category": "Thrift Store Consultant", "month": 1, "amount": 100000},
        {"category": "Equipment & Fixtures", "month": 1, "amount": 45000},
        {"category": "Build-out / Renovations / Tenant Improvements",
         "month": 1, "amount": 50000},
        {"category": "Legal, Permits, Insurance Setup, Signage",
         "month": 1, "amount": 13500},
    ]
    assert model.default_cost_schedule() == frozen
    assert model.default_cost_schedule(model.EXAMPLE_BUILDING_SQFT) == frozen


def test_default_cost_schedule_scales_only_the_two_derived_lines():
    sched = {row["category"]: row["amount"]
             for row in model.default_cost_schedule(20000)}
    # The two sq-ft lines double; consultant + legal stay flat.
    assert sched["Equipment & Fixtures"] == 90000
    assert sched["Build-out / Renovations / Tenant Improvements"] == 100000
    assert sched["Thrift Store Consultant"] == 100000
    assert sched["Legal, Permits, Insurance Setup, Signage"] == 13500


def test_default_cost_schedule_amounts_are_whole_dollars_at_10k():
    # Byte-stability guard: whole-dollar seeds stay ints (not 45000.0) so the
    # frozen migration literal + json.dumps stay identical.
    for row in model.default_cost_schedule():
        assert isinstance(row["amount"], int)


def test_default_cost_schedule_json_serializes_identically_at_10k():
    # The migration mirrors json.dumps(default_cost_schedule()); at 10k it must be
    # byte-identical to the frozen literal (ints, not 45000.0).
    import json
    dumped = json.dumps(model.default_cost_schedule())
    assert "45000" in dumped and "45000.0" not in dumped
    assert "50000" in dumped and "50000.0" not in dumped


def test_default_cost_schedule_non_integral_size_keeps_cents():
    # An odd sq ft that yields a fractional dollar is preserved (float), never
    # silently truncated to the wrong whole dollar.
    sched = {row["category"]: row["amount"] for row in model.default_cost_schedule(8333)}
    # 4.50 * 8333 = 37,498.5 ; 5.00 * 8333 = 41,665 (whole).
    assert sched["Equipment & Fixtures"] == 37498.5
    assert sched["Build-out / Renovations / Tenant Improvements"] == 41665


def test_utilities_is_flat_across_years_not_escalated():
    # Sq ft drives utilities, not the rent escalator — so it is identical Y1=Y2=Y3
    # (unlike Rent, which compounds). Use a non-10k size to make the value distinct.
    res = compute(_inputs("12500"), Assumptions(), DemandDrivers(), _trade())
    vals = [D(str(res.annual[y].expenses["Utilities/trash service"])) for y in (1, 2, 3)]
    assert vals[0] == vals[1] == vals[2] == Decimal("2.23") * Decimal("12500")
    # Rent, by contrast, DOES escalate — guards against accidentally coupling them.
    assert D(str(res.annual[2].expenses["Rent"])) > D(str(res.annual[1].expenses["Rent"]))


def test_no_schedule_startup_lump_unchanged_at_10k():
    # The no-cost_schedule path (Assumptions.startup lump) is NOT the scaled seed; at
    # the example basis it stays at today's 45,000 / 50,000 literals.
    res = compute(_inputs("10000"), Assumptions(), DemandDrivers(), _trade())
    assert D(str(res.startup["Equipment & Fixtures"])) == Decimal("45000")
    assert D(str(res.startup["Build-out / Renovations / Tenant Improvements"])) == Decimal("50000")


def test_viewmodel_rent_per_sqft_surfaced_when_input_provided():
    from engine.viewmodel import to_view_model
    inputs = Inputs(
        org_name="HoH", city="x", opening_status="pre_launch",
        share_y1="0.06", share_y2="0.08", share_y3="0.10",
        share_y4="0.11", share_y5="0.12", donation_value="74",
        building_sqft="10000", salesfloor_pct="0.70", rent_per_sqft="12",
        lease_type="nnn", source_of_funds="300000")
    vm = to_view_model(compute(inputs, Assumptions(), DemandDrivers(), _trade()))
    assert vm["settings"]["rentPerSqft"] == 12.0
    assert vm["settings"]["utilitiesPerSqft"] == 2.23
    # rent_per_sqft × sqft resolves occupancy (parity with a $120k flat entry).
    res = compute(inputs, Assumptions(), DemandDrivers(), _trade())
    assert D(str(res.occupancy[1])) == Decimal("120000")


def test_viewmodel_rent_per_sqft_null_when_flat_occupancy():
    from engine.viewmodel import to_view_model
    vm = to_view_model(compute(_inputs("10000"), Assumptions(), DemandDrivers(), _trade()))
    assert vm["settings"]["rentPerSqft"] is None  # occupancy entered as a flat annual $


def test_xlsx_utilities_row_is_a_live_sqft_formula_not_a_literal(tmp_path):
    # The workbook must DERIVE utilities (building_sqft × utilities_per_sqft) as a live
    # formula over blue input cells, so a LibreOffice recalc foots to the engine value.
    import openpyxl

    from engine.xlsx import build_xlsx
    res = compute(_inputs("10000"), Assumptions(), DemandDrivers(), _trade())
    path = tmp_path / "u.xlsx"
    build_xlsx(res, str(path))
    ws = openpyxl.load_workbook(path)["Projection"]
    labels = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    # The two new blue input cells exist.
    assert "building_sqft" in labels
    assert "utilities_per_sqft" in labels
    urow = labels["Utilities/trash service"]
    formula = ws.cell(row=urow, column=2).value
    assert isinstance(formula, str) and formula.startswith("=")
    assert "*" in formula  # building_sqft * utilities_per_sqft, not a hard literal
