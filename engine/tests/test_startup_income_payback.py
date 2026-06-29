"""RED first (epic #321, Phase 1): Year-1 income = full startup hit + real payback.

These tests assert the NEW, owner-locked contract and FAIL against current code
(compute.py / decision.py / monthly.py unchanged). They are the executable spec for
Phase 2; do NOT make them pass by touching engine logic in Phase 1.

Two owner decisions are pinned here:

1. INCOME = full Year-1 startup hit. The *reported* Year-1 net income subtracts the
   one-time startup total: ``reported_net[1] = total_rev[1] − total_exp[1] −
   startup_total``. Years 2–5 stay operating net. The view-model "Net income" KPI /
   row / ``incomeSeries[0]`` and the ``.xlsx`` "Net income" Year-1 cell reflect the
   hit. The *operating* net (``annual[y].net``) is preserved unchanged so the
   decision metrics, the monthly roll-up, the sensitivity re-runs and — by owner
   recommendation — the displayed Net **margin** and the §12.2 margin guardrail stay
   on operating economics (a one-time capital outlay must not masquerade as a margin).

2. PAYBACK = real "Month N" from grand opening. The monthly model is extended beyond
   the 12 Year-1 operating months across the full 5-year horizon, so payback resolves
   to the first project-month whose **cumulative cash ≥ 0** (startup fully recovered),
   instead of being structurally N/A (the Aurora bug). Capital required, break-even
   month and first-year cash gap stay **Year-1-anchored and byte-unchanged**.
"""
import shutil
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from engine.money import D
from engine.viewmodel import to_view_model
from engine.xlsx import build_xlsx


def _kpi(vm, label):
    return next(k for k in vm["kpis"] if k["label"] == label)


def _row(vm, label):
    return next(r for r in vm["rows"] if r["label"] == label)


# ---------------------------------------------------------------- (a) startup hit
def test_year1_net_income_includes_full_startup_hit(design_parity_result):
    # RED: today the "Net income" surfaces show the OPERATING Year-1 net; the contract
    # requires the *reported* Year-1 net to drop by exactly startup_total.
    res = design_parity_result
    vm = to_view_model(res)

    operating_net_y1 = D(str(res.annual[1].net))      # operating (kept, unchanged)
    startup_total = D(str(res.startup_total))
    assert startup_total > 0
    expected_reported = operating_net_y1 - startup_total   # the new Year-1 income

    # KPI strip, income sparkline series, and the Net income table row all reflect it.
    assert D(str(_kpi(vm, "Net income")["value"]["raw"])) == expected_reported
    assert D(str(vm["incomeSeries"][0]["raw"])) == expected_reported
    assert D(str(_row(vm, "Net income")["y1"]["raw"])) == expected_reported


def test_years_2_to_5_net_income_stays_operating(design_parity_result):
    # Years 2–5 are unchanged operating net (the startup hit lands only in Year 1).
    res = design_parity_result
    vm = to_view_model(res)
    ni_row = _row(vm, "Net income")
    for y, key in ((2, "y2"), (3, "y3"), (4, "y4"), (5, "y5")):
        assert D(str(ni_row[key]["raw"])) == D(str(res.annual[y].net))
        assert D(str(vm["incomeSeries"][y - 1]["raw"])) == D(str(res.annual[y].net))
    # year5Income is operating == reported for Y5 (unchanged).
    assert D(str(vm["year5Income"]["raw"])) == D(str(res.annual[5].net))


def test_year1_drop_equals_startup_total_for_canonical(canonical_result):
    # The drop is exactly startup_total for the loss fixture too (an even deeper Y1).
    res = canonical_result
    vm = to_view_model(res)
    expected = D(str(res.annual[1].net)) - D(str(res.startup_total))
    assert D(str(_kpi(vm, "Net income")["value"]["raw"])) == expected
    assert D(str(vm["incomeSeries"][0]["raw"])) == expected


# ------------------------------------------------------ (b) payback resolves (Aurora)
def test_positive_operating_net_yields_finite_payback_month(design_parity_result):
    # RED (the Aurora-class case): a clearly viable "proceed" projection with a POSITIVE
    # Year-1 operating net but startup > one year's net. Today payback_month is None
    # because the monthly model walks only 12 months; the extended horizon must resolve
    # it to a real "Month N" beyond Year 1.
    res = design_parity_result
    assert res.annual[1].net > 0                              # positive operating net
    assert D(str(res.startup_total)) > D(str(res.annual[1].net))  # 1yr can't recover it

    pm = res.decision.payback_month
    assert pm is not None, "payback must resolve to a real month, not N/A"
    assert 12 < pm <= 60, f"payback lands past Year 1 within the horizon, got {pm}"

    # Cash-basis definition: cumulative cash is < 0 before recovery and ≥ 0 at payback.
    by_month = {m.index: D(str(m.cumulative_cash)) for m in res.monthly}
    assert by_month[pm] >= 0
    assert any(c < 0 for i, c in by_month.items() if i < pm)

    # The web milestone and the Payback-period KPI all read the SAME month.
    vm = to_view_model(res)
    pb = vm["milestones"]["payback"]
    assert pb["raw"] == pm and pb["display"] == f"Month {pm}"
    assert _kpi(vm, "Payback period")["value"]["raw"] == pm


def test_canonical_loss_payback_stays_na_within_horizon(canonical_result):
    # The deep-loss fixture never recovers startup even over the full horizon, so its
    # payback stays explicit N/A-with-reason (never a misleading 0) — the extension
    # must not paper over a genuinely unrecovered case.
    res = canonical_result
    assert res.decision.payback_month is None
    assert res.flags.get("payback_month", "").startswith("N/A")


# ---------------------------------------------- (c) extended monthly roll-up (Reg #15)
def test_annual_rolls_up_from_extended_monthly_series(design_parity_result):
    # RED: the monthly series now spans the full 5-year horizon (60 operating months,
    # opening at month 1). Year-1 annual P&L still == roll-up of its 12 operating months
    # (Regression #15), and each later year rolls up to its own annual operating totals.
    res = design_parity_result
    ops = [m for m in res.monthly if m.kind == "operating"]
    assert len(ops) == 60, f"expected 60 operating months across the horizon, got {len(ops)}"

    for y in (1, 2, 3, 4, 5):
        seg = ops[(y - 1) * 12: y * 12]
        assert len(seg) == 12
        rev = sum((D(str(m.revenue)) for m in seg), Decimal("0"))
        exp = sum((D(str(m.expenses)) for m in seg), Decimal("0"))
        net = sum((D(str(m.net)) for m in seg), Decimal("0"))
        assert rev == D(str(res.annual[y].total_rev)), f"Y{y} revenue roll-up"
        assert exp == D(str(res.annual[y].total_exp)), f"Y{y} expense roll-up"
        assert net == D(str(res.annual[y].net)), f"Y{y} operating-net roll-up"


# --------------------------------------------------- (d) xlsx Net income Y1 cell
def test_xlsx_net_income_year1_subtracts_total_startup(design_parity_result, tmp_path):
    # RED (no LibreOffice needed): the formula-driven workbook must subtract the
    # "Total Startup Costs" cell from the Year-1 Net income formula, so the in-sheet
    # value foots to the server's startup-inclusive net. Today the formula is
    # ``=B{tr}-B{te}`` with no startup reference.
    res = design_parity_result
    xlsx = tmp_path / "p.xlsx"
    build_xlsx(res, str(xlsx))
    wb = load_workbook(str(xlsx))                      # formulas (not data_only)
    ws = wb["Projection"]
    label_row = {
        ws.cell(row=r, column=1).value: r
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=1).value, str)
    }
    ni_row = label_row["Net income"]
    st_row = label_row["Total Startup Costs"]
    ni_y1 = ws.cell(row=ni_row, column=2).value        # column B = Year 1
    wb.close()
    assert isinstance(ni_y1, str) and ni_y1.startswith("=")
    assert f"B{st_row}" in ni_y1, (
        f"Year-1 Net income formula {ni_y1!r} must subtract Total Startup Costs "
        f"cell B{st_row}")


@pytest.mark.skipif(shutil.which("soffice") is None,
                    reason="LibreOffice (soffice) not installed")
def test_xlsx_net_income_year1_recalcs_to_server_net(design_parity_result, tmp_path):
    # RED where LibreOffice exists: a real recalc of the Year-1 Net income cell must
    # equal the server's startup-inclusive net within $0.50 (the golden-gate tolerance).
    import pathlib
    import sys
    sys.path.insert(0, str(
        pathlib.Path(__file__).resolve().parents[2]
        / "reference/thrift-store-projections-skill/scripts"))
    import recalc as recalc_harness

    res = design_parity_result
    vm = to_view_model(res)
    xlsx = tmp_path / "p.xlsx"
    build_xlsx(res, str(xlsx))
    report = recalc_harness.recalc(str(xlsx), timeout=90)
    assert "error" not in report and report["total_errors"] == 0, report

    wb = load_workbook(str(xlsx), data_only=True)
    ws = wb["Projection"]
    label_row = {
        ws.cell(row=r, column=1).value: r
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=1).value, str)
    }
    cell = ws.cell(row=label_row["Net income"], column=2).value
    wb.close()
    expected = float(_row(vm, "Net income")["y1"]["raw"])
    assert abs(float(cell) - expected) <= 0.5


# --------------------------------------- invariant locks (GREEN now; must STAY green)
def test_net_margin_row_stays_operating_not_startup_inclusive(design_parity_result):
    # Owner recommendation: the displayed Net margin (and the §12.2 guardrail) stay on
    # OPERATING economics — margin Y1 = operating_net / revenue, NOT the startup-hit
    # net / revenue. This is true today and MUST remain true after Phase 2 (so the
    # Net income row and the Net margin row are intentionally decoupled in Year 1).
    res = design_parity_result
    vm = to_view_model(res)
    operating_margin_y1 = D(str(res.annual[1].net)) / D(str(res.annual[1].total_rev))
    assert abs(D(str(_row(vm, "Net margin")["y1"]["raw"])) - operating_margin_y1) < D("1e-9")


def test_capital_and_cash_gap_are_year1_window_anchored(canonical_result):
    # Capital required and first-year cash gap must be derivable from the Year-1
    # operating window (pre-opening + months G..G+11) ALONE, so extending the monthly
    # model across 5 years cannot move them (Year-1-anchored, byte-unchanged contract).
    res = canonical_result
    dec = res.decision
    g = dec.grand_opening_month
    window = [D(str(m.cumulative_cash)) for m in res.monthly if m.index <= g + 11]
    trough = min(window)
    assert D(str(dec.first_year_cash_gap)) == min(
        D(str(m.cumulative_cash)) for m in res.monthly if g <= m.index <= g + 11)
    assert D(str(dec.startup_capital_required)) == abs(min(Decimal("0"), trough)) + Decimal("25000")
    # break-even stays the Year-1 ramp month (the frozen canonical anchor).
    assert dec.break_even_month == 9
