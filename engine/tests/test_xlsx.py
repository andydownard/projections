"""RED first: the .xlsx cell contract (§11) — structure, guards, overrides.

The recalc/reconciliation against LibreOffice lives in the golden-gate test
(test_golden_gate.py); here we assert the workbook *structure*: blue inputs,
live formulas, totals as formulas, division guards, the monthly tab, and the
orange overridden-component cell + exact comment string.
"""

import openpyxl
import pytest

from engine.compute import Assumptions, DemandDrivers, Inputs, TradeArea, compute
from engine.xlsx import build_xlsx


def _inputs():
    return Inputs(
        org_name="House of Hope", city="Longview, Texas", opening_status="pre_launch",
        share_y1="0.06", share_y2="0.08", share_y3="0.10", share_y4="0.11", share_y5="0.12",
        donation_value="74", building_sqft="10000", salesfloor_pct="0.70",
        occupancy_annual="120000", lease_type="nnn", source_of_funds="300000",
    )


def _build(tmp_path, overrides=None):
    res = compute(_inputs(), Assumptions(), DemandDrivers(),
                  TradeArea(basis="city", households="31798"), overrides=overrides or {})
    path = tmp_path / "proj.xlsx"
    build_xlsx(res, str(path))
    return res, path


def test_two_tabs_projection_and_monthly(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    assert "Projection" in wb.sheetnames
    assert "Monthly Cash Flow" in wb.sheetnames


def test_totals_are_formulas_not_literals(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    formulas = [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=SUM(") for f in formulas)        # totals are SUMs
    assert any("/" in f and "IF(" in f for f in formulas)       # guarded divisions


def test_no_unguarded_division_by_a_cell(tmp_path):
    import re
    # A division whose denominator is a CELL reference (could be 0) must be
    # guarded. Division by a literal constant (e.g. /12) is always safe.
    cell_denom = re.compile(r"/\s*\$?[A-Z]{1,3}\$?\d+")
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if (isinstance(c.value, str) and c.value.startswith("=")
                        and cell_denom.search(c.value)):
                    assert "IF(" in c.value, \
                        f"unguarded cell division at {ws.title}!{c.coordinate}: {c.value}"


def test_year2_share_not_average(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert "AVERAGE(" not in c.value  # midpoint is dead


def test_overridden_cell_is_orange_with_exact_comment(tmp_path):
    overrides = {"Payroll#1": {"value": "200000", "note": "frozen", "ts": "2026-06-24"}}
    _res, path = _build(tmp_path, overrides=overrides)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    # find the Payroll Year-1 cell: hard-coded numeric 200000, orange, commented.
    found = None
    for row in ws.iter_rows():
        label = row[0].value
        if label == "Payroll":
            found = row[1]  # Year-1 column (B)
            break
    assert found is not None
    assert float(found.value) == 200000.0           # hard-coded, not a formula
    assert found.comment is not None
    assert found.comment.text.startswith("Manual override — calc was =")
    assert "set to $200,000 on 2026-06-24" in found.comment.text


def _projection_label_map(ws):
    """Map a column-A string label -> its row index on the Projection sheet."""
    return {
        ws.cell(row=r, column=1).value: r
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=1).value, str)
    }


def _all_text(ws):
    return " ".join(
        str(c.value) for row in ws.iter_rows() for c in row
        if isinstance(c.value, str)
    )


# ----------------------------------------------- #192 export disclaimer in the .xlsx
def test_projection_sheet_has_notes_and_disclaimer_block(tmp_path):
    # #192: the workbook a board reads detached from the app must carry the
    # "planning estimates, not advice" caveat — server-formatted text, no client math.
    from engine.model import EXPORT_DISCLAIMER
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)
    assert "NOTES & DISCLAIMER" in labels
    assert EXPORT_DISCLAIMER in _all_text(ws)


def test_xlsx_disclaimer_is_text_only_no_formula(tmp_path):
    # The disclaimer block is pure server-formatted text: no formulas (so it can
    # never introduce a recalc error onto the correctness spine).
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)
    start = labels["NOTES & DISCLAIMER"]
    for r in range(start, ws.max_row + 1):
        for col in range(1, 4):
            v = ws.cell(row=r, column=col).value
            assert not (isinstance(v, str) and v.startswith("=")), f"formula at row {r}"


def test_xlsx_disclaimer_stamps_terms_version_and_accepted_date(tmp_path):
    # When the caller knows the terms version + acceptance date they are stamped, so
    # the artifact records WHICH terms the operator accepted and WHEN (#192).
    from engine.model import TERMS_VERSION
    res = compute(_inputs(), Assumptions(), DemandDrivers(),
                  TradeArea(basis="city", households="31798"))
    path = tmp_path / "stamped.xlsx"
    build_xlsx(res, str(path), terms_accepted_at="2026-06-27T12:00:00Z")
    wb = openpyxl.load_workbook(path)
    text = _all_text(wb["Projection"])
    assert f"Terms v{TERMS_VERSION}" in text
    assert "2026-06-27" in text


def test_projection_sheet_has_milestones_block(tmp_path):
    # The workbook reflects the three milestones (§8 / #117) consistently with the
    # view-model: grand opening = G, break-even, payback — each as its server display.
    res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)
    assert "MILESTONES" in labels
    for name in ("Grand opening", "Break-even", "Payback"):
        assert name in labels, name
    from engine.viewmodel import to_view_model
    ms = to_view_model(res)["milestones"]
    assert ws.cell(row=labels["Grand opening"], column=2).value == ms["grandOpening"]["display"]
    assert ws.cell(row=labels["Break-even"], column=2).value == ms["breakEven"]["display"]
    assert ws.cell(row=labels["Payback"], column=2).value == ms["payback"]["display"]


def test_milestones_block_renders_payback_na_not_zero(tmp_path):
    # The canonical default never recovers startup within the horizon: the workbook
    # must show an explicit N/A, never a misleading 0 (matches the view-model).
    res, path = _build(tmp_path)
    assert res.decision.payback_month is None
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)
    assert ws.cell(row=labels["Payback"], column=2).value == "N/A"


def test_milestones_block_shows_payback_month_when_recovered(tmp_path):
    # A strongly profitable case recovers startup within the operating ramp; the
    # workbook shows the payback month (consistent with the view-model display).
    inputs = Inputs(
        org_name="Strong", city="x", opening_status="pre_launch",
        share_y1="0.15", share_y2="0.20", share_y3="0.25", share_y4="0.11", share_y5="0.12",
        donation_value="74", building_sqft="10000", salesfloor_pct="0.70",
        occupancy_annual="120000", lease_type="nnn", source_of_funds="800000",
    )
    res = compute(inputs, Assumptions(labor_pct="0.40"), DemandDrivers(),
                  TradeArea(basis="city", households="31798"))
    assert res.decision.payback_month is not None
    path = tmp_path / "strong.xlsx"
    build_xlsx(res, str(path))
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)
    assert ws.cell(row=labels["Payback"], column=2).value == f"Month {res.decision.payback_month}"


def _inputs_with_cost_schedule(schedule):
    return Inputs(
        org_name="House of Hope", city="Longview, Texas", opening_status="pre_launch",
        share_y1="0.06", share_y2="0.08", share_y3="0.10", share_y4="0.11", share_y5="0.12",
        donation_value="74", building_sqft="10000", salesfloor_pct="0.70",
        occupancy_annual="120000", lease_type="nnn", source_of_funds="300000",
        cost_schedule=schedule, opening_month_number=1,
    )


def test_startup_separate_and_contingent_sections_render_none_when_empty(tmp_path):
    # Model v3 Phase 2 (Task 2.2): the canonical fixture has no cost_schedule, so
    # both new sub-sections are ALWAYS rendered but empty -- "None", never omitted,
    # so a reader never wonders whether the section was silently dropped.
    res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)
    sep_row = labels["Funded separately (not included in capital required)"]
    cont_row = labels["Contingent (not included in capital required)"]
    assert ws.cell(row=sep_row + 1, column=1).value == "None"
    assert ws.cell(row=cont_row + 1, column=1).value == "None"
    # placed after the (included-only) Total Startup Costs line, before MILESTONES.
    assert labels["Total Startup Costs"] < sep_row < cont_row < labels["MILESTONES"]


def test_startup_separate_and_contingent_sections_list_rows_and_subtotal(tmp_path):
    schedule = (
        {"category": "Included line", "month": 1, "amount": 10000,
         "treatment": "included"},
        {"category": "Consultant", "month": 1, "amount": 100000,
         "treatment": "separate"},
        {"category": "Contingency reserve", "month": 1, "amount": 20000,
         "treatment": "contingent"},
    )
    res = compute(_inputs_with_cost_schedule(schedule), Assumptions(), DemandDrivers(),
                  TradeArea(basis="city", households="31798"), overrides={})
    path = tmp_path / "sep_cont.xlsx"
    build_xlsx(res, str(path))
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)

    total_row = labels["Total Startup Costs"]
    sep_row = labels["Funded separately (not included in capital required)"]
    cont_row = labels["Contingent (not included in capital required)"]
    consultant_row = labels["Consultant"]
    contingency_row = labels["Contingency reserve"]

    # both new sections -- and the rows inside them -- sit AFTER the (included
    # -only) Total Startup Costs line, so neither can be inside its =SUM(...) range.
    assert total_row < sep_row < cont_row
    assert consultant_row > total_row
    assert contingency_row > total_row

    assert ws.cell(row=consultant_row, column=2).value == 100000.0
    assert ws.cell(row=consultant_row + 1, column=1).value == "Subtotal"
    assert (ws.cell(row=consultant_row + 1, column=2).value
            == f"=SUM(B{consultant_row}:B{consultant_row})")

    assert ws.cell(row=contingency_row, column=2).value == 20000.0
    assert ws.cell(row=contingency_row + 1, column=1).value == "Subtotal"
    assert (ws.cell(row=contingency_row + 1, column=2).value
            == f"=SUM(B{contingency_row}:B{contingency_row})")

    # the (included-only) Total Startup Costs total is still a SUM formula over
    # its own component rows (proven numerically by the golden-gate recalc).
    assert ws.cell(row=total_row, column=2).value.startswith("=SUM(")


def test_monthly_tab_has_twelve_month_rows(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    months = [c.value for col in ws.iter_cols(min_col=1, max_col=1) for c in col
              if isinstance(c.value, str) and c.value.startswith("Month ")]
    assert len(months) >= 12


# --------------------------------------------------------------- pre-opening
def _preopening_inputs():
    """A projection with grand opening on month 3 and a timed cost schedule that
    places costs on the two pre-opening months and the opening month."""
    return Inputs(
        org_name="House of Hope", city="Longview, Texas", opening_status="pre_launch",
        share_y1="0.06", share_y2="0.08", share_y3="0.10", share_y4="0.11", share_y5="0.12",
        donation_value="74", building_sqft="10000", salesfloor_pct="0.70",
        occupancy_annual="120000", lease_type="nnn", source_of_funds="300000",
        opening_month_number=3,
        cost_schedule=(
            {"category": "Consultant", "month": 1, "amount": "20000"},
            {"category": "Build-out", "month": 2, "amount": "50000"},
            {"category": "Equipment", "month": 3, "amount": "30000"},
        ),
    )


def _build_preopening(tmp_path, overrides=None):
    res = compute(_preopening_inputs(), Assumptions(), DemandDrivers(),
                  TradeArea(basis="city", households="31798"), overrides=overrides or {})
    path = tmp_path / "preopen.xlsx"
    build_xlsx(res, str(path))
    return res, path


def _month_row_map(ws):
    """Map project-month int -> worksheet row index, parsing the 'Month N' label."""
    import re
    out = {}
    for col in ws.iter_cols(min_col=1, max_col=1):
        for c in col:
            if isinstance(c.value, str):
                m = re.match(r"Month (\d+)", c.value)
                if m:
                    out[int(m.group(1))] = c.row
    return out


def test_monthly_tab_emits_a_row_per_project_month_1_to_n(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    rowmap = _month_row_map(ws)
    expected = [m.index for m in res.monthly]
    assert sorted(rowmap) == expected
    assert expected[0] == 1
    assert len(expected) == res.inputs.opening_month_number + 59  # N = G + 59


def test_monthly_tab_has_scheduled_costs_column(tmp_path):
    _res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = [ws.cell(row=2, column=col).value for col in range(1, 7)]
    assert "Scheduled costs" in headers


def test_scheduled_costs_land_on_their_month_numbers(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = [ws.cell(row=2, column=col).value for col in range(1, 7)]
    sched_col = headers.index("Scheduled costs") + 1
    rowmap = _month_row_map(ws)
    by_month = {m.index: m.startup_outflow for m in res.monthly}
    for month, r in rowmap.items():
        cell = ws.cell(row=r, column=sched_col).value
        assert float(cell) == float(by_month[month]), f"month {month}"
    # the two explicit pre-opening schedule rows landed on months 1 and 2
    assert float(ws.cell(row=rowmap[1], column=sched_col).value) == 20000.0
    assert float(ws.cell(row=rowmap[2], column=sched_col).value) == 50000.0


def test_pre_opening_months_have_negative_net_from_scheduled_costs(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = [ws.cell(row=2, column=col).value for col in range(1, 7)]
    rev_col = headers.index("Revenue") + 1
    exp_col = headers.index("Operating expenses") + 1
    sched_col = headers.index("Scheduled costs") + 1
    net_col = headers.index("Net") + 1
    rowmap = _month_row_map(ws)
    pre_open = [m.index for m in res.monthly if m.kind == "pre_open"]
    assert pre_open == [1, 2]
    for month in pre_open:
        r = rowmap[month]
        # no operating activity
        assert float(ws.cell(row=r, column=rev_col).value) == 0.0
        assert float(ws.cell(row=r, column=exp_col).value) == 0.0
        # net is a live formula over revenue - opex - scheduled costs (so it is
        # the negative scheduled cost — a real monthly loss)
        net = ws.cell(row=r, column=net_col).value
        assert isinstance(net, str) and net.startswith("=")
        from openpyxl.utils import get_column_letter
        rc = get_column_letter(rev_col)
        ec = get_column_letter(exp_col)
        sc = get_column_letter(sched_col)
        assert net == f"={rc}{r}-{ec}{r}-{sc}{r}"


def test_preopening_workbook_has_no_unguarded_divisions(tmp_path):
    import re
    cell_denom = re.compile(r"/\s*\$?[A-Z]{1,3}\$?\d+")
    _res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if (isinstance(c.value, str) and c.value.startswith("=")
                        and cell_denom.search(c.value)):
                    assert "IF(" in c.value, \
                        f"unguarded cell division at {ws.title}!{c.coordinate}"


def _headers(ws):
    return [ws.cell(row=2, column=col).value for col in range(1, 7)]


def test_monthly_headers_are_in_exact_order(tmp_path):
    _res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    assert _headers(ws) == [
        "Month", "Revenue", "Operating expenses", "Scheduled costs",
        "Net", "Cumulative cash",
    ]


def test_monthly_revenue_and_expenses_are_literals_not_formulas(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = _headers(ws)
    rev_col = headers.index("Revenue") + 1
    exp_col = headers.index("Operating expenses") + 1
    rowmap = _month_row_map(ws)
    for r in rowmap.values():
        rev = ws.cell(row=r, column=rev_col).value
        exp = ws.cell(row=r, column=exp_col).value
        assert not (isinstance(rev, str) and rev.startswith("="))
        assert not (isinstance(exp, str) and exp.startswith("="))


def test_operating_months_carry_revenue_and_expenses(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = _headers(ws)
    rev_col = headers.index("Revenue") + 1
    exp_col = headers.index("Operating expenses") + 1
    rowmap = _month_row_map(ws)
    by_month = {m.index: m for m in res.monthly}
    operating = [m.index for m in res.monthly
                 if m.kind == "operating" and m.index <= res.inputs.opening_month_number + 11]
    assert operating  # there are operating months
    for month in operating:
        r = rowmap[month]
        # approx, not ==: the workbook round-trips a float and the eased months carry
        # long Decimal expansions (sub-millionth-of-a-cent differences are real but
        # irrelevant; the golden gate itself uses a $0.50 tolerance).
        assert float(ws.cell(row=r, column=rev_col).value) == pytest.approx(
            float(by_month[month].revenue), abs=1e-6)
        assert float(ws.cell(row=r, column=exp_col).value) == pytest.approx(
            float(by_month[month].expenses), abs=1e-6)
    # at least one operating month has positive revenue (the ramp is non-trivial)
    assert any(float(by_month[m].revenue) > 0 for m in operating)


def test_monthly_net_is_formula_for_every_month(tmp_path):
    from openpyxl.utils import get_column_letter
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = _headers(ws)
    rev = get_column_letter(headers.index("Revenue") + 1)
    exp = get_column_letter(headers.index("Operating expenses") + 1)
    sched = get_column_letter(headers.index("Scheduled costs") + 1)
    net_col = headers.index("Net") + 1
    rowmap = _month_row_map(ws)
    for r in rowmap.values():
        assert ws.cell(row=r, column=net_col).value == f"={rev}{r}-{exp}{r}-{sched}{r}"


def test_cumulative_cash_is_running_sum_of_net(tmp_path):
    from openpyxl.utils import get_column_letter
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = _headers(ws)
    net = get_column_letter(headers.index("Net") + 1)
    cum_col = headers.index("Cumulative cash") + 1
    cum = get_column_letter(cum_col)
    rowmap = _month_row_map(ws)
    ordered = [rowmap[m] for m in sorted(rowmap)]
    first = ordered[0]
    assert ws.cell(row=first, column=cum_col).value == f"={net}{first}"
    for prev, r in zip(ordered, ordered[1:]):
        assert ws.cell(row=r, column=cum_col).value == f"={cum}{prev}+{net}{r}"


def test_scheduled_costs_sum_equals_startup_total(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = _headers(ws)
    sched_col = headers.index("Scheduled costs") + 1
    rowmap = _month_row_map(ws)
    total = sum(float(ws.cell(row=r, column=sched_col).value) for r in rowmap.values())
    assert abs(total - float(res.startup_total)) < 0.5


def test_opening_month_scheduled_cost_includes_lease_deposit(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = _headers(ws)
    sched_col = headers.index("Scheduled costs") + 1
    rowmap = _month_row_map(ws)
    g = res.inputs.opening_month_number
    # lease deposit = occupancy_y1 / 12 * 3 = 120000/12*3 = 30000, plus the
    # explicit month-G equipment cost of 30000 => 60000 on the opening month.
    assert float(ws.cell(row=rowmap[g], column=sched_col).value) == 60000.0


def test_opening_month_label_is_bold_milestone(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    rowmap = _month_row_map(ws)
    g = res.inputs.opening_month_number
    cell = ws.cell(row=rowmap[g], column=1)
    assert "Grand opening" in cell.value
    assert cell.font.bold


def test_monthly_cells_use_usd_number_format(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    rowmap = _month_row_map(ws)
    for r in rowmap.values():
        for col in range(2, 7):
            assert "$" in ws.cell(row=r, column=col).number_format


def test_scheduled_costs_are_never_negative(tmp_path):
    res, path = _build_preopening(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = _headers(ws)
    sched_col = headers.index("Scheduled costs") + 1
    rowmap = _month_row_map(ws)
    for r in rowmap.values():
        assert float(ws.cell(row=r, column=sched_col).value) >= 0.0


def test_default_no_schedule_has_no_pre_opening_rows(tmp_path):
    # The default inputs open on month 1, so there are no pre-opening months; the tab
    # spans the full 60-month horizon (CASH model 2.7.0).
    res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    rowmap = _month_row_map(ws)
    assert sorted(rowmap) == list(range(1, 61))
    assert all(m.kind == "operating" for m in res.monthly)


def test_default_no_schedule_places_whole_lump_on_opening_month(tmp_path):
    res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = _headers(ws)
    sched_col = headers.index("Scheduled costs") + 1
    rowmap = _month_row_map(ws)
    # opening is month 1; the entire startup total lands there.
    assert abs(float(ws.cell(row=rowmap[1], column=sched_col).value)
               - float(res.startup_total)) < 0.5
    # every other month (across the full 60-month horizon) carries no scheduled cost.
    for month in range(2, 61):
        assert float(ws.cell(row=rowmap[month], column=sched_col).value) == 0.0


def test_post_open_scheduled_cost_extends_the_timeline(tmp_path):
    # A support cost beyond the 60-month operating window must still get a row on its
    # own month number; N extends to cover it (N = max(G+59, last scheduled month)).
    inputs = Inputs(
        org_name="House of Hope", city="Longview, Texas", opening_status="pre_launch",
        share_y1="0.06", share_y2="0.08", share_y3="0.10", share_y4="0.11", share_y5="0.12",
        donation_value="74", building_sqft="10000", salesfloor_pct="0.70",
        occupancy_annual="120000", lease_type="nnn", source_of_funds="300000",
        opening_month_number=2,
        cost_schedule=(
            {"category": "Consultant", "month": 1, "amount": "15000"},
            {"category": "Post-open support", "month": 64, "amount": "9000"},
        ),
    )
    res = compute(inputs, Assumptions(), DemandDrivers(),
                  TradeArea(basis="city", households="31798"))
    path = tmp_path / "postopen.xlsx"
    build_xlsx(res, str(path))
    wb = openpyxl.load_workbook(path)
    ws = wb["Monthly Cash Flow"]
    headers = _headers(ws)
    sched_col = headers.index("Scheduled costs") + 1
    rowmap = _month_row_map(ws)
    assert max(rowmap) == 64
    assert sorted(rowmap) == list(range(1, 65))
    assert float(ws.cell(row=rowmap[64], column=sched_col).value) == 9000.0


def test_expense_row_skipped_when_label_absent_from_annual(tmp_path):
    # build_xlsx walks EXPENSE_ORDER and `continue`s past any label the year's
    # expenses dict doesn't carry (xlsx.py). Drop one fixed-annual line so its
    # EXPENSE_ORDER label is genuinely absent, and assert the workbook simply omits
    # that row (no crash, no phantom row) while a full build still renders it.
    from engine import model
    trimmed = {k: v for k, v in model.FIXED_ANNUAL.items() if k != "Pest control"}
    res = compute(_inputs(), Assumptions(fixed_annual=trimmed), DemandDrivers(),
                  TradeArea(basis="city", households="31798"))
    assert "Pest control" not in res.annual[1].expenses  # the absent label
    path = tmp_path / "trimmed.xlsx"
    build_xlsx(res, str(path))
    ws = openpyxl.load_workbook(path)["Projection"]
    assert "Pest control" not in _projection_label_map(ws)  # row skipped
    # a present line is unaffected — still rendered.
    assert "Utilities/trash service" in _projection_label_map(ws)

    # Contrast: the default build (all fixed lines present) DOES render the row.
    full = compute(_inputs(), Assumptions(), DemandDrivers(),
                   TradeArea(basis="city", households="31798"))
    full_path = tmp_path / "full.xlsx"
    build_xlsx(full, str(full_path))
    full_ws = openpyxl.load_workbook(full_path)["Projection"]
    assert "Pest control" in _projection_label_map(full_ws)


# ---------------------------------------- Task 1.4: newly-editable-setting inputs
def _label_rows(ws, label):
    """Every row whose column-A value equals ``label``. A fixed-annual line appears
    twice — once as a KEY-INPUTS blue cell, once as a 5-YEAR-PROJECTION expense row —
    so ``_projection_label_map`` (last-occurrence) can't distinguish the two."""
    return [r for r in range(1, ws.max_row + 1)
            if ws.cell(row=r, column=1).value == label]


# ---------------------------------------- Task 1.2: production-station module
def test_production_station_rows_are_guarded_formulas(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)
    for label in ("Items sold", "Items produced", "Production hours"):
        assert label in labels, label
        for col in range(2, 7):
            v = ws.cell(row=labels[label], column=col).value
            assert isinstance(v, str) and v.startswith("=IF("), (label, v)


def test_production_station_input_cells_are_blue(tmp_path):
    _res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)
    for name in ("avg_item_price", "sell_through", "items_per_hour",
                 "productive_hours_per_station"):
        assert name in labels, name
        cell = ws.cell(row=labels[name], column=2)
        assert not (isinstance(cell.value, str) and cell.value.startswith("="))
        assert cell.font.color.rgb.endswith("0000FF"), f"{name} input not blue"


def test_production_stations_needed_cell_near_milestones(tmp_path):
    res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    labels = _projection_label_map(ws)
    assert "Production stations needed (peak month)" in labels
    cell = ws.cell(row=labels["Production stations needed (peak month)"], column=2)
    # server-resolved plain number, not a formula (matches how milestone strings
    # are server snapshots, not live divisions).
    assert not (isinstance(cell.value, str) and cell.value.startswith("="))
    assert cell.value == res.stations_installed


def test_fixed_annual_lines_are_blue_inputs_referenced_by_formula(tmp_path):
    """Task 1.4 contract: each now-editable ``Assumptions.fixed_annual`` line (e.g.
    'Pest control') is a BLUE input cell in KEY INPUTS, and its 5-YEAR-PROJECTION
    expense row is a live formula referencing that input cell by coordinate (=$B$NN)
    across all five years — never a re-inlined literal. So editing the blue cell
    re-foots the expense row, Total Expenses (=SUM over it), Net income, and margin.

    This locks the fixed_annual wiring (the one newly-editable setting that feeds a
    sheet formula) against a future refactor silently collapsing it back to a literal.
    Reserve buffer / capital and the sensitivity delta have NO cell in this workbook
    (capital & sensitivity are engine-only decision metrics, absent from the sheet), so
    per scope they get no input cell here — see task report.
    """
    res, path = _build(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Projection"]
    assert res.assumptions.fixed_annual, "there are fixed-annual lines to check"
    for label in res.assumptions.fixed_annual:
        rows = _label_rows(ws, label)
        assert len(rows) == 2, f"{label}: expected a KEY-INPUTS row and an expense row"
        input_row, expense_row = rows  # KEY INPUTS precedes the 5-YEAR PROJECTION
        input_cell = ws.cell(row=input_row, column=2)
        # blue input cell carrying the literal annual $ (a number, not a formula)
        assert isinstance(input_cell.value, (int, float)), f"{label} input not numeric"
        assert input_cell.font.color.rgb.endswith("0000FF"), f"{label} input not blue"
        # each of Year 1..5 references the SINGLE input cell by ref — not a literal
        for col in range(2, 7):
            f = ws.cell(row=expense_row, column=col).value
            assert f == f"=$B${input_row}", f"{label} Y{col - 1} not ref to input: {f!r}"
