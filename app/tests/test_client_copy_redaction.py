"""CI redaction gate for client-copy export artifacts (Model v3 Phase 3, Task 3.2).

Proves that the client-copy `.xlsx` and packet markdown variants never contain
internal vocabulary (§12.2 citations, DAG jargon, flag keys, version strings, etc.)
that should be invisible to board/funder readers. Also verifies the internal variant
DOES contain these terms today, so the test harness itself is non-trivial and would
have caught a real leak.
"""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.packet import build_packet, render_packet_md
from engine.tests.conftest import build_aggressive_result
from engine.xlsx import build_client_copy_xlsx

BANNED_TERMS = [
    "§", "DAG", "override", "990-grounded", "benchmark store",
    "model_version", "compute_schema_version", "MODEL_VERSION", "COMPUTE_SCHEMA_VERSION",
    "net_margin_out_of_band", "labor_ratio_out_of_band", "share_y1_out_of_band",
    "occupancy_ratio_out_of_band", "sales_per_sqft_out_of_band",
    "leadership_coverage_advisory", "ramp_low_confidence",
    "LABOR_RATIO_BAND", "SHARE_Y1_BAND_MAX", "NET_MARGIN_FLAG_MAX",
    "engine/", "app/routers/", ".py:",
]


def _assert_no_banned_terms(text: str) -> None:
    """Verify none of the banned terms (case-insensitive) appear in the text.

    Raises AssertionError if any banned term is found.
    """
    text_lower = text.lower()
    for term in BANNED_TERMS:
        term_lower = term.lower()
        if term_lower in text_lower:
            raise AssertionError(
                f"Banned term '{term}' (case-insensitive) found in artifact text"
            )


def test_client_copy_xlsx_redaction():
    """Client-copy xlsx contains no banned internal vocabulary."""
    result = build_aggressive_result()
    # Build in memory
    path = BytesIO()
    build_client_copy_xlsx(result, path)
    path.seek(0)

    # Read back the workbook
    wb = load_workbook(path)

    # Collect all cell values and comments
    all_text = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    all_text.append(str(cell.value))
                if cell.comment:
                    all_text.append(str(cell.comment.text))

    combined_text = " ".join(all_text)
    _assert_no_banned_terms(combined_text)


def test_client_copy_packet_redaction():
    """Client-copy packet markdown contains no banned internal vocabulary."""
    result = build_aggressive_result()
    vm = result.view_model

    # Build packet with client variant
    pkt = build_packet(vm, role="board", variant="client")

    # Render to markdown
    markdown = render_packet_md(pkt)

    _assert_no_banned_terms(markdown)


def test_internal_packet_contains_banned_terms_proves_test_is_nontrivial():
    """Internal packet DOES contain banned terms — proves the test would catch a real leak.

    This is the negative-path test that documents the harness is non-trivial and
    would have caught the pre-existing (§12.2) citation leak in the internal variant.
    """
    result = build_aggressive_result()
    vm = result.view_model

    # Build packet with internal variant (default) — should have guardrails section
    pkt = build_packet(vm, role="board", variant="internal")

    # Render to markdown
    markdown = render_packet_md(pkt)

    # Verify the test harness is non-trivial: internal variant SHOULD contain
    # at least one banned term (the "§12.2" citation in guardrails).
    with pytest.raises(AssertionError):
        _assert_no_banned_terms(markdown)
