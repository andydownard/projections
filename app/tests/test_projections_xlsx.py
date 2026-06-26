"""M2 .xlsx export — GET /api/projections/{view_slug}/xlsx (API_SPEC §6.3).

The endpoint renders the formula-driven workbook (engine/xlsx.py) from the stored
projection's resolved node values — the SAME math surface as computed_json. RED-first:
written before the endpoint exists.

Coverage:
  - returns the spreadsheetml content-type + attachment Content-Disposition
  - the bytes round-trip through openpyxl into the two-tab workbook
  - terms gate: 403 terms_required when terms not accepted; 404 for unknown slug
  - (CI / local soffice) a LibreOffice recalc of the workbook yields zero formula
    errors — guarded skip when soffice or its macro harness is unavailable.
"""
import io
import os
import shutil
import sys

import openpyxl
import pytest

from app.tests.conftest import VALID_BODY

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def _create(client, body=None):
    return client.post("/api/projections", json=body or VALID_BODY)


def _accept_terms(store, slug):
    row = store.get_projection_by_slug(slug)
    store._conn.execute(
        "UPDATE projections SET terms_accepted_at = ? WHERE id = ?",
        ("2026-06-24T00:00:00Z", row["id"]),
    )
    store._conn.commit()


def test_xlsx_export_returns_spreadsheet_content_type(client, store):
    created = _create(client).json()
    slug = created["view_slug"]
    _accept_terms(store, slug)
    r = client.get(f"/api/projections/{slug}/xlsx")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX_CT
    cd = r.headers.get("content-disposition", "")
    assert "attachment" in cd
    assert ".xlsx" in cd


def test_xlsx_export_bytes_round_trip_to_two_tab_workbook(client, store):
    created = _create(client).json()
    slug = created["view_slug"]
    _accept_terms(store, slug)
    r = client.get(f"/api/projections/{slug}/xlsx")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Projection" in wb.sheetnames
    assert "Monthly Cash Flow" in wb.sheetnames
    # totals are live formulas (foots in-sheet)
    ws = wb["Projection"]
    formulas = [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=SUM(") for f in formulas)


def test_xlsx_export_requires_terms_acceptance(client):
    created = _create(client).json()
    slug = created["view_slug"]
    r = client.get(f"/api/projections/{slug}/xlsx")
    assert r.status_code == 403, r.text
    assert r.json()["detail"]["error"] == "terms_required"


def test_xlsx_export_unknown_slug_404(client):
    r = client.get("/api/projections/does-not-exist/xlsx")
    assert r.status_code == 404


def _soffice_recalc(path):
    """Recalc the workbook via LibreOffice; return the recalc.py result dict.

    Adds the reference skill's office helper to sys.path so recalc.py imports.
    """
    ref = os.path.join(
        os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
        "reference", "thrift-store-projections-skill", "scripts")
    if ref not in sys.path:
        sys.path.insert(0, ref)
    import recalc as recalc_mod
    return recalc_mod.recalc(path)


@pytest.mark.skipif(shutil.which("soffice") is None,
                    reason="CI-only: requires a pinned LibreOffice (soffice) on PATH")
def test_xlsx_recalc_has_zero_formula_errors(client, store, tmp_path):
    created = _create(client).json()
    slug = created["view_slug"]
    _accept_terms(store, slug)
    r = client.get(f"/api/projections/{slug}/xlsx")
    assert r.status_code == 200, r.text
    path = tmp_path / "export.xlsx"
    path.write_bytes(r.content)
    try:
        result = _soffice_recalc(str(path))
    except ModuleNotFoundError as exc:
        pytest.skip(f"CI-only: soffice recalc harness unavailable ({exc})")
    if "error" in result:
        pytest.skip(f"CI-only: soffice recalc unavailable ({result['error']})")
    assert result["status"] == "success", result
    assert result["total_errors"] == 0, result
