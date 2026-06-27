"""
Excel Formula Recalculation Script
Recalculates all formulas in an Excel file using LibreOffice
"""

import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from office.soffice import get_soffice_env

from openpyxl import load_workbook

MACRO_DIR_MACOS = "~/Library/Application Support/LibreOffice/4/user/basic/Standard"
MACRO_DIR_LINUX = "~/.config/libreoffice/4/user/basic/Standard"
MACRO_FILENAME = "Module1.xba"

RECALCULATE_MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""

# Registers Module1 in the Basic "Standard" library so the macro resolves via
# location=application even on a freshly-created per-invocation profile (#199).
SCRIPT_XLB = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE library:library PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "library.dtd">
<library:library xmlns:library="http://openoffice.org/2000/library" \
library:name="Standard" library:readonly="false" library:passwordprotected="false">
 <library:element library:name="Module1"/>
</library:library>"""


def _user_installation_arg(user_installation):
    """The ``-env:UserInstallation`` switch pointing soffice at an isolated profile."""
    return f"-env:UserInstallation={Path(user_installation).absolute().as_uri()}"


def _profile_macro_dir(user_installation):
    """The Basic ``Standard`` library dir for a profile.

    With ``user_installation`` set, it lives inside that throwaway profile; with
    ``None`` it is the shared host profile (legacy / direct-CLI behaviour).
    """
    if user_installation is None:
        return os.path.expanduser(
            MACRO_DIR_MACOS if platform.system() == "Darwin" else MACRO_DIR_LINUX
        )
    return os.path.join(user_installation, "user", "basic", "Standard")


def has_gtimeout():
    try:
        subprocess.run(
            ["gtimeout", "--version"], capture_output=True, timeout=1, check=False
        )
        return True
    except (FileNotFoundError, subprocess.TimeoutExpired):
        return False


def setup_libreoffice_macro(user_installation=None):
    """Install the recalc macro into a profile's Basic ``Standard`` library.

    With ``user_installation`` set, the macro is installed into that throwaway,
    per-invocation profile (hermetic, never touches the developer's real profile);
    with ``None`` it falls back to the shared host profile (legacy behaviour). (#199)
    """
    macro_dir = _profile_macro_dir(user_installation)
    macro_file = os.path.join(macro_dir, MACRO_FILENAME)

    if (
        os.path.exists(macro_file)
        and "RecalculateAndSave" in Path(macro_file).read_text()
    ):
        return True

    if not os.path.exists(macro_dir):
        init_cmd = ["soffice", "--headless", "--terminate_after_init"]
        if user_installation is not None:
            # Initialise THIS profile (creates the Standard library scaffolding) so
            # the recalc that follows reuses it instead of the shared default.
            init_cmd.insert(1, _user_installation_arg(user_installation))
        subprocess.run(
            init_cmd,
            capture_output=True,
            timeout=30,
            env=get_soffice_env(),
        )
        os.makedirs(macro_dir, exist_ok=True)

    try:
        Path(macro_file).write_text(RECALCULATE_MACRO)
        # Ensure the Standard library registers Module1 even on a hand-built profile.
        Path(os.path.join(macro_dir, "script.xlb")).write_text(SCRIPT_XLB)
        return True
    except Exception:
        return False


def recalc(filename, timeout=30, user_installation=None):
    """Recalculate ``filename`` in LibreOffice and report formula errors.

    Each call runs in its OWN throwaway LibreOffice profile (a unique
    ``-env:UserInstallation``) unless ``user_installation`` is supplied. LibreOffice
    is single-instance *per profile*: without isolation a concurrent call forwards
    the document to the already-running instance and the launcher returns 0 WITHOUT
    recalculating — yielding empty cached cells and a flaky red on the golden gate.
    A per-invocation profile makes recalc hermetic (safe under pytest-xdist) and
    stops polluting the developer's real profile. (#199)
    """
    if not Path(filename).exists():
        return {"error": f"File {filename} does not exist"}

    abs_path = str(Path(filename).absolute())

    own_profile = user_installation is None
    if own_profile:
        user_installation = tempfile.mkdtemp(prefix="lo_recalc_profile_")

    try:
        if not setup_libreoffice_macro(user_installation):
            return {"error": "Failed to setup LibreOffice macro"}

        cmd = [
            "soffice",
            _user_installation_arg(user_installation),
            "--headless",
            "--norestore",
            "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
            abs_path,
        ]

        if platform.system() == "Linux":
            cmd = ["timeout", str(timeout)] + cmd
        elif platform.system() == "Darwin" and has_gtimeout():
            cmd = ["gtimeout", str(timeout)] + cmd

        return _run_recalc(cmd, filename)
    finally:
        if own_profile:
            shutil.rmtree(user_installation, ignore_errors=True)


def _run_recalc(cmd, filename):
    result = subprocess.run(cmd, capture_output=True, text=True, env=get_soffice_env())

    if result.returncode == 124:
        # gtimeout/timeout killed soffice mid-recalc — fail closed, never load a
        # partially/un-recalculated workbook and report it as a success.
        return {"error": "recalc timed out"}
    if result.returncode != 0:
        error_msg = result.stderr or "Unknown error during recalculation"
        if "Module1" in error_msg or "RecalculateAndSave" not in error_msg:
            return {"error": "LibreOffice macro not configured properly"}
        return {"error": error_msg}

    try:
        wb = load_workbook(filename, data_only=True)

        excel_errors = [
            "#VALUE!",
            "#DIV/0!",
            "#REF!",
            "#NAME?",
            "#NULL!",
            "#NUM!",
            "#N/A",
        ]
        error_details = {err: [] for err in excel_errors}
        total_errors = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None and isinstance(cell.value, str):
                        for err in excel_errors:
                            if err in cell.value:
                                location = f"{sheet_name}!{cell.coordinate}"
                                error_details[err].append(location)
                                total_errors += 1
                                break

        wb.close()

        result = {
            "status": "success" if total_errors == 0 else "errors_found",
            "total_errors": total_errors,
            "error_summary": {},
        }

        for err_type, locations in error_details.items():
            if locations:
                result["error_summary"][err_type] = {
                    "count": len(locations),
                    "locations": locations[:20],  
                }

        wb_formulas = load_workbook(filename, data_only=False)
        formula_count = 0
        for sheet_name in wb_formulas.sheetnames:
            ws = wb_formulas[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if (
                        cell.value
                        and isinstance(cell.value, str)
                        and cell.value.startswith("=")
                    ):
                        formula_count += 1
        wb_formulas.close()

        result["total_formulas"] = formula_count

        return result

    except Exception as e:
        return {"error": str(e)}


def main():
    if len(sys.argv) < 2:
        print("Usage: python recalc.py <excel_file> [timeout_seconds]")
        print("\nRecalculates all formulas in an Excel file using LibreOffice")
        print("\nReturns JSON with error details:")
        print("  - status: 'success' or 'errors_found'")
        print("  - total_errors: Total number of Excel errors found")
        print("  - total_formulas: Number of formulas in the file")
        print("  - error_summary: Breakdown by error type with locations")
        print("    - #VALUE!, #DIV/0!, #REF!, #NAME?, #NULL!, #NUM!, #N/A")
        sys.exit(1)

    filename = sys.argv[1]
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 30

    result = recalc(filename, timeout)
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
