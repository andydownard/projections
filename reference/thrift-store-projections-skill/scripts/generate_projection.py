#!/usr/bin/env python3
"""
Thrift Store 3-Year Projection generator.

Bottom-up model driven by donor market share:
  donating households = city households x donor market share
  annual donations    = donating households x donations per household per year (default 3)
  retail revenue      = annual donations x average value per donation
  customers           = retail revenue / average transaction value
Expenses keep the standard thrift line items; rent = sq ft x rent per sq ft;
payroll = labor % of revenue.

Produces:
  1) <Org> 3-Year Projection.xlsx  -- live, formula-driven, edit the blue input cells
  2) <Org> 3-Year Projection.docx  -- one-page summary, simple to ingest

Usage:
  python generate_projection.py config.json /path/to/output_dir
"""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------- defaults
DEFAULTS = {
    "donate_freq": 3,           # average household donates 3x / year
    "salvage_pct": 0.008,       # salvage sales as % of retail
    "cc_pct": 0.02,             # credit card fees as % of revenue
    "labor_pct": 0.60,          # payroll as % of revenue
    "payroll_tax_pct": 0.075,   # payroll taxes as % of payroll
    "fixed": {
        "Insurance liability": 2000,
        "Internet, Phones": 3000,
        "Pest control": 300,
        "Repairs/maintenance": 1600,
        "Security system": 900,
        "Store/office supplies": 3700,
        "Utilities/trash service": 22300,
    },
    "startup": {
        "Thrift Store Consultant": 100000,
        "Equipment & Fixtures": 45000,
        "Build-out / Renovations / Tenant Improvements": 50000,
        "Legal, Permits, Insurance Setup, Signage": 13500,
    },
}

def merge(cfg):
    c = dict(DEFAULTS)
    c.update({k: v for k, v in cfg.items() if k not in ("fixed", "startup")})
    c["fixed"] = dict(DEFAULTS["fixed"]); c["fixed"].update(cfg.get("fixed", {}))
    c["startup"] = dict(DEFAULTS["startup"]); c["startup"].update(cfg.get("startup", {}))
    return c

# ---------------------------------------------------------------- compute (mirrors Excel)
def compute(c):
    H, freq = c["households"], c["donate_freq"]
    share = {1: c["share_y1"], 2: (c["share_y1"] + c["share_y3"]) / 2, 3: c["share_y3"]}
    r = {y: {} for y in (1, 2, 3)}
    rent = c["sqft"] * c["rent_per_sqft"]
    for y in (1, 2, 3):
        d = r[y]
        d["share"] = share[y]
        d["donating_hh"] = H * share[y]
        d["donations"] = d["donating_hh"] * freq
        d["retail"] = d["donations"] * c["donation_value"]
        d["customers"] = d["retail"] / c["transaction_value"]
        d["salvage"] = d["retail"] * c["salvage_pct"]
        d["total_rev"] = d["retail"] + d["salvage"]
        d["Credit card fees"] = d["total_rev"] * c["cc_pct"]
        d["Payroll"] = d["total_rev"] * c["labor_pct"]
        d["Payroll taxes"] = d["Payroll"] * c["payroll_tax_pct"]
        d["Rent"] = rent
        for k, v in c["fixed"].items():
            d[k] = v
        exp_lines = ["Credit card fees", "Insurance liability", "Internet, Phones",
                     "Payroll", "Payroll taxes", "Pest control", "Rent",
                     "Repairs/maintenance", "Security system", "Store/office supplies",
                     "Utilities/trash service"]
        d["total_exp"] = sum(d[k] for k in exp_lines)
        d["net"] = d["total_rev"] - d["total_exp"]
        d["margin"] = d["net"] / d["total_rev"] if d["total_rev"] else 0
    startup = dict(c["startup"])
    startup["Lease Deposit + First Month Rent"] = rent / 12 * 3
    startup_total = sum(startup.values())
    payback = startup_total / r[1]["net"] if r[1]["net"] else 0
    return r, rent, startup, startup_total, payback

EXP_LINES = ["Credit card fees", "Insurance liability", "Internet, Phones",
             "Payroll", "Payroll taxes", "Pest control", "Rent",
             "Repairs/maintenance", "Security system", "Store/office supplies",
             "Utilities/trash service"]

# ---------------------------------------------------------------- Excel
BLUE = Font(name="Arial", color="0000FF", size=10)
BLACK = Font(name="Arial", color="000000", size=10)
BOLD = Font(name="Arial", bold=True, size=10)
WHITEHDR = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", fgColor="1F4E5F")
SUB_FILL = PatternFill("solid", fgColor="D9E7EC")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
TOT_FILL = PatternFill("solid", fgColor="EAF1F4")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%'
NUM = '#,##0'

def style_input(cell, fmt=None):
    cell.font = BLUE; cell.fill = INPUT_FILL; cell.border = BORDER
    if fmt: cell.number_format = fmt
    cell.alignment = Alignment(horizontal="right")

def build_xlsx(c, path):
    r, rent, startup, startup_total, payback = compute(c)
    wb = Workbook(); ws = wb.active; ws.title = "Projection"
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 40, "B": 15, "C": 15, "D": 15, "E": 34}.items():
        ws.column_dimensions[col].width = w

    def hdr(row, text):
        ws.merge_cells(f"A{row}:D{row}")
        cl = ws[f"A{row}"]; cl.value = text; cl.font = WHITEHDR; cl.fill = HDR_FILL
        cl.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 18

    # Title
    ws.merge_cells("A1:D1")
    ws["A1"] = f"{c['org_name']} — 3-Year Thrift Store Projection"
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws.merge_cells("A2:E2")
    ws["A2"] = f"{c['city']}  ·  Bottom-up model driven by donor market share  ·  Edit the blue cells to update every projection"
    ws["A2"].font = Font(name="Arial", italic=True, size=9, color="555555")

    # ---- INPUTS
    hdr(4, "KEY INPUTS  ·  edit the blue cells")
    inputs = [  # (label, value, fmt, note, attr)
        ("Households in market (city)", c["households"], NUM, "Source: U.S. Census (households in city)", "households"),
        ("Donations per household per year", c["donate_freq"], NUM, "Average household donates ~3x / year", "donate_freq"),
        ("Donor market share — Year 1", c["share_y1"], PCT, "Share of households that donate to this store", "share_y1"),
        ("Donor market share — Year 3", c["share_y3"], PCT, "Year 2 is the midpoint of Year 1 and Year 3", "share_y3"),
        ("Average value per donation ($)", c["donation_value"], USD, "Resale value generated per donation → top-line revenue", "donation_value"),
        ("Average transaction value ($)", c["transaction_value"], USD, "Average customer basket → number of customers", "transaction_value"),
        ("Store size (sq ft)", c["sqft"], NUM, "Drives rent", "sqft"),
        ("Rent ($/sq ft/year)", c["rent_per_sqft"], USD, "Local market rate", "rent_per_sqft"),
        ("Labor (% of revenue)", c["labor_pct"], PCT, "Payroll as a share of revenue", "labor_pct"),
        ("Payroll tax (% of payroll)", c["payroll_tax_pct"], PCT, "", "payroll_tax_pct"),
        ("Credit card fees (% of revenue)", c["cc_pct"], PCT, "", "cc_pct"),
        ("Salvage sales (% of retail)", c["salvage_pct"], PCT, "Recycled/salvage income", "salvage_pct"),
    ]
    cell_ref = {}
    row = 5
    for label, value, fmt, note, attr in inputs:
        ws[f"A{row}"] = label; ws[f"A{row}"].font = BLACK
        cl = ws[f"B{row}"]; cl.value = value; style_input(cl, fmt)
        if note:
            ws[f"E{row}"] = note; ws[f"E{row}"].font = Font(name="Arial", size=8, color="808080")
        cell_ref[attr] = f"$B${row}"
        row += 1

    row += 0
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = "Fixed annual operating costs"; ws[f"A{row}"].font = BOLD; ws[f"A{row}"].fill = SUB_FILL
    row += 1
    fixed_ref = {}
    for label, value in c["fixed"].items():
        ws[f"A{row}"] = label; ws[f"A{row}"].font = BLACK
        cl = ws[f"B{row}"]; cl.value = value; style_input(cl, USD)
        fixed_ref[label] = f"$B${row}"
        row += 1

    # ---- PROJECTION
    row += 1
    proj_hdr = row
    hdr(proj_hdr, "3-YEAR PROJECTION  ·  calculated from the inputs above")
    row += 1
    yrs = ["Year 1", "Year 2", "Year 3"]
    ws[f"A{row}"] = "Operating drivers"; ws[f"A{row}"].font = BOLD
    for i, y in enumerate(yrs):
        cl = ws.cell(row=row, column=2 + i, value=y); cl.font = BOLD
        cl.alignment = Alignment(horizontal="right"); cl.fill = SUB_FILL
    row += 1
    cols = ["B", "C", "D"]

    def line(label, formulas, fmt=USD, bold=False, fill=None, top=False):
        nonlocal row
        a = ws[f"A{row}"]; a.value = label
        a.font = BOLD if bold else BLACK
        for i, f in enumerate(formulas):
            cl = ws.cell(row=row, column=2 + i, value=f)
            cl.font = BOLD if bold else BLACK
            cl.number_format = fmt
            cl.alignment = Alignment(horizontal="right")
            if fill: cl.fill = fill
            if top: cl.border = Border(top=Side(style="thin", color="808080"))
        if fill: a.fill = fill
        row += 1

    # share row: Y1 input, Y2 midpoint, Y3 input
    share_row = row
    line("Donor market share", [f"={cell_ref['share_y1']}",
                                 f"=AVERAGE({cell_ref['share_y1']},{cell_ref['share_y3']})",
                                 f"={cell_ref['share_y3']}"], PCT)
    dh_row = row
    line("Donating households",
         [f"={cell_ref['households']}*{c}{share_row}" for c in cols], NUM)
    don_row = row
    line("Annual donations",
         [f"={c}{dh_row}*{cell_ref['donate_freq']}" for c in cols], NUM)
    # retail defined below; customers references retail row
    cust_row = row
    line("Customers", ["", "", ""], NUM)  # placeholder, fill after retail known

    row += 1
    ws[f"A{row}"] = "Income statement"; ws[f"A{row}"].font = BOLD
    for i in range(3):
        ws.cell(row=row, column=2 + i).fill = SUB_FILL
    row += 1
    retail_row = row
    line("Retail Sales", [f"={c}{don_row}*{cell_ref['donation_value']}" for c in cols], USD)
    salvage_row = row
    line("Salvage Sales", [f"={c}{retail_row}*{cell_ref['salvage_pct']}" for c in cols], USD)
    totrev_row = row
    line("Total Revenue", [f"={c}{retail_row}+{c}{salvage_row}" for c in cols], USD, bold=True, top=True)
    # now fill customers = retail / transaction value
    for i, c_ in enumerate(cols):
        cl = ws.cell(row=cust_row, column=2 + i, value=f"={c_}{retail_row}/{cell_ref['transaction_value']}")
        cl.number_format = NUM; cl.font = BLACK; cl.alignment = Alignment(horizontal="right")

    row += 0
    # expense lines
    exp_first = row
    line("Credit card fees", [f"={c}{totrev_row}*{cell_ref['cc_pct']}" for c in cols], USD)
    line("Insurance liability", [f"={fixed_ref['Insurance liability']}"] * 3, USD)
    line("Internet, Phones", [f"={fixed_ref['Internet, Phones']}"] * 3, USD)
    payroll_row = row
    line("Payroll", [f"={c}{totrev_row}*{cell_ref['labor_pct']}" for c in cols], USD)
    line("Payroll taxes", [f"={c}{payroll_row}*{cell_ref['payroll_tax_pct']}" for c in cols], USD)
    line("Pest control", [f"={fixed_ref['Pest control']}"] * 3, USD)
    line("Rent", [f"={cell_ref['sqft']}*{cell_ref['rent_per_sqft']}"] * 3, USD)
    line("Repairs/maintenance", [f"={fixed_ref['Repairs/maintenance']}"] * 3, USD)
    line("Security system", [f"={fixed_ref['Security system']}"] * 3, USD)
    line("Store/office supplies", [f"={fixed_ref['Store/office supplies']}"] * 3, USD)
    line("Utilities/trash service", [f"={fixed_ref['Utilities/trash service']}"] * 3, USD)
    exp_last = row - 1
    totexp_row = row
    line("Total Expenses", [f"=SUM({c}{exp_first}:{c}{exp_last})" for c in cols], USD, bold=True, top=True)
    net_row = row
    line("Net profit", [f"={c}{totrev_row}-{c}{totexp_row}" for c in cols], USD, bold=True, fill=TOT_FILL)
    line("Net margin", [f"={c}{net_row}/{c}{totrev_row}" for c in cols], PCT)

    # ---- STARTUP
    row += 1
    hdr(row, "ONE-TIME STARTUP COSTS  ·  Year 0 (capital, not in operating expenses)")
    row += 1
    su_first = row
    order = ["Thrift Store Consultant", "Equipment & Fixtures",
             "Lease Deposit + First Month Rent",
             "Build-out / Renovations / Tenant Improvements",
             "Legal, Permits, Insurance Setup, Signage"]
    for k in order:
        ws[f"A{row}"] = k; ws[f"A{row}"].font = BLACK
        if k == "Lease Deposit + First Month Rent":
            cl = ws[f"B{row}"]; cl.value = f"={cell_ref['sqft']}*{cell_ref['rent_per_sqft']}/12*3"
            cl.font = BLACK
        else:
            cl = ws[f"B{row}"]; cl.value = c["startup"][k]; style_input(cl, USD)
        cl.number_format = USD; cl.alignment = Alignment(horizontal="right")
        row += 1
    su_last = row - 1
    ws[f"A{row}"] = "Total Startup Costs"; ws[f"A{row}"].font = BOLD
    cl = ws[f"B{row}"]; cl.value = f"=SUM(B{su_first}:B{su_last})"; cl.font = BOLD
    cl.number_format = USD; cl.fill = TOT_FILL; cl.alignment = Alignment(horizontal="right")
    startup_total_row = row
    row += 1
    ws[f"A{row}"] = "Simple payback (years)"; ws[f"A{row}"].font = BLACK
    cl = ws[f"B{row}"]; cl.value = f"=B{startup_total_row}/B{net_row}"; cl.font = BLACK
    cl.number_format = '0.0'; cl.alignment = Alignment(horizontal="right")
    ws[f"E{row}"] = "Total startup ÷ Year 1 net profit"; ws[f"E{row}"].font = Font(name="Arial", size=8, color="808080")

    wb.save(path)

# ---------------------------------------------------------------- Word (1 page)
def _set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def _no_space(p):
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)

def money(x): return f"${x:,.0f}"

def build_docx(c, path):
    r, rent, startup, startup_total, payback = compute(c)
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(9)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    sec.left_margin = sec.right_margin = Inches(0.6)

    NAVY = RGBColor(0x1F, 0x4E, 0x5F)
    title = doc.add_paragraph(); _no_space(title)
    run = title.add_run(f"{c['org_name']} — 3-Year Thrift Store Projection")
    run.bold = True; run.font.size = Pt(15); run.font.color.rgb = NAVY
    sub = doc.add_paragraph(); _no_space(sub)
    sr = sub.add_run(f"{c['city']}  ·  Bottom-up model driven by donor market share")
    sr.italic = True; sr.font.size = Pt(9); sr.font.color.rgb = RGBColor(0x66,0x66,0x66)
    sub.paragraph_format.space_after = Pt(6)

    y1 = r[1]
    summ = doc.add_paragraph(); _no_space(summ)
    summ.add_run("Year 1 outlook: ").bold = True
    summ.add_run(
        f"{y1['donations']:,.0f} donations and {y1['customers']:,.0f} customers, "
        f"generating {money(y1['total_rev'])} in revenue and {money(y1['net'])} net profit "
        f"({y1['margin']*100:.0f}% margin). By Year 3, revenue reaches {money(r[3]['total_rev'])} "
        f"on a {c['share_y3']*100:.0f}% donor market share.")
    summ.paragraph_format.space_after = Pt(6)

    # Assumptions strip
    ap = doc.add_paragraph(); _no_space(ap)
    ap.add_run("Key assumptions  ").bold = True
    assume = (f"{c['households']:,} city households · ~{c['donate_freq']}x donations/household/yr · "
              f"market share {c['share_y1']*100:.0f}% → { (c['share_y1']+c['share_y3'])/2*100:.0f}% → {c['share_y3']*100:.0f}% (Yr1→3) · "
              f"{money(c['donation_value'])}/donation · {money(c['transaction_value'])} avg basket · "
              f"{c['sqft']:,} sq ft @ {money(c['rent_per_sqft'])}/sq ft · labor {c['labor_pct']*100:.0f}% of revenue")
    ar = ap.add_run(assume); ar.font.size = Pt(8.5); ar.font.color.rgb = RGBColor(0x44,0x44,0x44)
    ap.paragraph_format.space_after = Pt(6)

    # Income statement table
    rows_def = [("Operating drivers", None, "sub")]
    rows_def += [("Donor market share", [f"{r[y]['share']*100:.0f}%" for y in (1,2,3)], "drv")]
    rows_def += [("Annual donations", [f"{r[y]['donations']:,.0f}" for y in (1,2,3)], "drv")]
    rows_def += [("Customers", [f"{r[y]['customers']:,.0f}" for y in (1,2,3)], "drv")]
    rows_def += [("Revenue", None, "sub")]
    rows_def += [("Retail Sales", [money(r[y]['retail']) for y in (1,2,3)], "row")]
    rows_def += [("Salvage Sales", [money(r[y]['salvage']) for y in (1,2,3)], "row")]
    rows_def += [("Total Revenue", [money(r[y]['total_rev']) for y in (1,2,3)], "tot")]
    rows_def += [("Expenses", None, "sub")]
    for k in EXP_LINES:
        rows_def += [(k, [money(r[y][k]) for y in (1,2,3)], "row")]
    rows_def += [("Total Expenses", [money(r[y]['total_exp']) for y in (1,2,3)], "tot")]
    rows_def += [("Net profit", [money(r[y]['net']) for y in (1,2,3)], "net")]
    rows_def += [("Net margin", [f"{r[y]['margin']*100:.0f}%" for y in (1,2,3)], "row")]

    table = doc.add_table(rows=1 + len(rows_def), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(3.3), Inches(1.4), Inches(1.4), Inches(1.4)]
    # header
    hdr_cells = table.rows[0].cells
    headers = ["Projected Income Statement", "Year 1", "Year 2", "Year 3"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]; _no_space(p)
        run = p.add_run(h); run.bold = True; run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        _set_cell_bg(hdr_cells[i], "1F4E5F")
    for ri, (label, vals, kind) in enumerate(rows_def, start=1):
        cells = table.rows[ri].cells
        if vals is None:  # subheader spanning
            a = cells[0].merge(cells[3])
            p = a.paragraphs[0]; _no_space(p)
            run = p.add_run(label); run.bold = True; run.font.size = Pt(8.5)
            _set_cell_bg(a, "D9E7EC")
            continue
        p0 = cells[0].paragraphs[0]; _no_space(p0)
        run0 = p0.add_run(("   " if kind in ("row","drv") else "") + label)
        run0.font.size = Pt(9); run0.bold = kind in ("tot","net")
        for i, v in enumerate(vals):
            pc = cells[i+1].paragraphs[0]; _no_space(pc)
            pc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            rr = pc.add_run(v); rr.font.size = Pt(9); rr.bold = kind in ("tot","net")
            if kind == "net": rr.font.color.rgb = NAVY
        if kind == "net":
            for cc in cells: _set_cell_bg(cc, "EAF1F4")
    for rrow in table.rows:
        for i, cell in enumerate(rrow.cells):
            cell.width = widths[i]
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    _thin_table_borders(table)

    # Startup costs compact line
    sp = doc.add_paragraph(); _no_space(sp)
    sp.paragraph_format.space_before = Pt(6)
    sp.add_run("One-time startup costs (Year 0):  ").bold = True
    parts = []
    order = ["Thrift Store Consultant", "Equipment & Fixtures",
             "Lease Deposit + First Month Rent",
             "Build-out / Renovations / Tenant Improvements",
             "Legal, Permits, Insurance Setup, Signage"]
    for k in order:
        parts.append(f"{k} {money(startup[k])}")
    sr2 = sp.add_run("  ·  ".join(parts)); sr2.font.size = Pt(8.5)
    tp = doc.add_paragraph(); _no_space(tp)
    tr = tp.add_run(f"Total startup {money(startup_total)}  ·  simple payback {payback:.1f} years "
                    f"(total startup ÷ Year 1 net profit).")
    tr.bold = True; tr.font.size = Pt(9)

    fn = doc.add_paragraph(); _no_space(fn); fn.paragraph_format.space_before = Pt(8)
    fr = fn.add_run("Method: donating households = city households × market share; annual donations = "
                    "donating households × 3; retail revenue = donations × value per donation; "
                    "customers = revenue ÷ average basket. Edit any assumption in the accompanying spreadsheet.")
    fr.italic = True; fr.font.size = Pt(7.5); fr.font.color.rgb = RGBColor(0x80,0x80,0x80)

    doc.save(path)

def _thin_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'BFBFBF')
        borders.append(e)
    tblPr.append(borders)

# ---------------------------------------------------------------- main
def main():
    if len(sys.argv) < 3:
        print("usage: generate_projection.py config.json output_dir"); sys.exit(1)
    cfg = merge(json.load(open(sys.argv[1])))
    outdir = sys.argv[2]; os.makedirs(outdir, exist_ok=True)
    base = f"{cfg['org_name']} 3-Year Projection"
    xlsx = os.path.join(outdir, base + ".xlsx")
    docx = os.path.join(outdir, base + ".docx")
    build_xlsx(cfg, xlsx)
    build_docx(cfg, docx)
    print(json.dumps({"xlsx": xlsx, "docx": docx}))

if __name__ == "__main__":
    main()
